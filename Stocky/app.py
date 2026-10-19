import json
import os
import sqlite3
import smtplib
import threading
import sys
from datetime import datetime
from email.mime.text import MIMEText

import tkinter as tk
from tkinter import ttk, messagebox, simpledialog

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
except ImportError:
    openpyxl = None

try:
    import win32print
    import win32ui
    import win32con
    HAS_WIN32PRINT = True
except ImportError:
    HAS_WIN32PRINT = False

try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


# ----------------------------------------------------------------------
# Rutas y constantes
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "stock.db")
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
VENTAS_XLSX = os.path.join(BASE_DIR, "ventas.xlsx")
STOCK_XLSX = os.path.join(BASE_DIR, "stock.xlsx")
STOCK_TXT = os.path.join(BASE_DIR, "stock.txt")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo.jpg")

DEFAULT_CONFIG = {
    "tema": "claro",              # "claro" u "oscuro"
    "tamano_fuente": 11,
    "email_alertas": "",          # a dónde se manda el aviso de stock bajo
    "email_remitente": "",        # cuenta de Gmail que envía el aviso
    "email_password": "",         # contraseña de aplicación de Gmail
    "smtp_server": "smtp.gmail.com",
    "smtp_port": 465,
    "umbral_stock_bajo": 3,
    "nombre_impresora": "",       # vacío = usar la impresora predeterminada
    "nombre_local": "MI COMERCIO",
    "mensaje_ticket": "¡Gracias por su compra!",
    "disposicion_ticket": "Centro", # "Izquierda", "Centro", "Derecha"
}

VENTAS_HEADERS = ["Día", "Fecha", "Hora", "Producto", "Precio Unitario", "Cantidad", "Subtotal", "Ganancia"]


# ----------------------------------------------------------------------
# Configuración (persistida en config.json)
# ----------------------------------------------------------------------
def cargar_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            cfg = DEFAULT_CONFIG.copy()
            cfg.update(data)
            return cfg
        except Exception:
            pass
    return DEFAULT_CONFIG.copy()


def guardar_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)


# ----------------------------------------------------------------------
# Base de datos (SQLite) - productos / stock
# ----------------------------------------------------------------------
def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS productos (
            codigo TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            precio REAL NOT NULL,
            stock INTEGER NOT NULL DEFAULT 0
        )
        """
    )
    return conn


def buscar_producto(codigo):
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos WHERE codigo = ?", (codigo,))
    row = cur.fetchone()
    conn.close()
    return row  # (codigo, nombre, precio, stock) o None


def buscar_producto_por_nombre(nombre):
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos WHERE nombre = ?", (nombre,))
    row = cur.fetchone()
    conn.close()
    return row


def guardar_producto(codigo, nombre, precio, stock):
    conn = get_conn()
    conn.execute(
        "INSERT INTO productos (codigo, nombre, precio, stock) VALUES (?,?,?,?) "
        "ON CONFLICT(codigo) DO UPDATE SET nombre=excluded.nombre, precio=excluded.precio, stock=excluded.stock",
        (codigo, nombre, precio, stock),
    )
    conn.commit()
    conn.close()


def eliminar_producto(codigo):
    conn = get_conn()
    conn.execute("DELETE FROM productos WHERE codigo = ?", (codigo,))
    conn.commit()
    conn.close()


def listar_productos():
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos ORDER BY nombre ASC")
    rows = cur.fetchall()
    conn.close()
    return rows


def descontar_stock(codigo, cantidad):
    """Descuenta stock y devuelve (nombre, stock_restante)."""
    conn = get_conn()
    cur = conn.execute("SELECT nombre, stock FROM productos WHERE codigo = ?", (codigo,))
    row = cur.fetchone()
    if row is None:
        conn.close()
        return None, None
    nombre, stock_actual = row
    nuevo_stock = max(0, stock_actual - cantidad)
    conn.execute("UPDATE productos SET stock = ? WHERE codigo = ?", (nuevo_stock, codigo))
    conn.commit()
    conn.close()
    return nombre, nuevo_stock


# ----------------------------------------------------------------------
# Email de alerta de stock bajo
# ----------------------------------------------------------------------
def enviar_alerta_stock(cfg, nombre_producto, stock_restante):
    """Envía el email en un hilo aparte para no trabar la interfaz."""
    destino = cfg.get("email_alertas", "").strip()
    remitente = cfg.get("email_remitente", "").strip()
    password = cfg.get("email_password", "").strip()
    if not destino or not remitente or not password:
        return  # alertas no configuradas, se ignora silenciosamente

    def _enviar():
        try:
            asunto = f"Stock bajo: {nombre_producto}"
            cuerpo = (
                f"El producto '{nombre_producto}' quedó con {stock_restante} unidad(es) "
                f"en stock (umbral configurado: {cfg.get('umbral_stock_bajo')})."
            )
            msg = MIMEText(cuerpo, "plain", "utf-8")
            msg["Subject"] = asunto
            msg["From"] = remitente
            msg["To"] = destino
            with smtplib.SMTP_SSL(cfg.get("smtp_server", "smtp.gmail.com"), int(cfg.get("smtp_port", 465))) as server:
                server.login(remitente, password)
                server.sendmail(remitente, [destino], msg.as_string())
        except Exception as e:
            print(f"[Aviso] No se pudo enviar el email de alerta: {e}")

    threading.Thread(target=_enviar, daemon=True).start()


# ----------------------------------------------------------------------
# Impresión térmica y normal
# ----------------------------------------------------------------------
ESC = b"\x1b"
GS = b"\x1d"
CORTE_PAPEL = GS + b"V" + b"\x00"      # corta el papel
INIT_PRINTER = ESC + b"@"              # reset impresora
CENTRAR = ESC + b"a" + b"\x01"
ALINEAR_IZQ = ESC + b"a" + b"\x00"
ALINEAR_DER = ESC + b"a" + b"\x02"
NEGRITA_ON = ESC + b"E" + b"\x01"
NEGRITA_OFF = ESC + b"E" + b"\x00"


def obtener_comando_alineacion(disposicion):
    if disposicion == "Izquierda":
        return ALINEAR_IZQ
    elif disposicion == "Derecha":
        return ALINEAR_DER
    return CENTRAR


def construir_ticket(cfg, items, subtotal, margen, total_final):
    """items: lista de dicts {codigo, nombre, precio, cantidad}"""
    lineas = []
    lineas.append(INIT_PRINTER)
    
    alineacion = obtener_comando_alineacion(cfg.get("disposicion_ticket", "Centro"))
    lineas.append(alineacion)
    
    nombre_local = cfg.get("nombre_local", "MI COMERCIO").upper()
    lineas.append(NEGRITA_ON + f"{nombre_local}\n".encode("cp437", "replace") + NEGRITA_OFF)
    lineas.append(b"TICKET DE VENTA\n")
    lineas.append(ALINEAR_IZQ)
    ahora = datetime.now()
    lineas.append(f"{ahora.strftime('%d/%m/%Y')}  {ahora.strftime('%H:%M:%S')}\n".encode("cp437", "replace"))
    lineas.append(b"--------------------------------\n")
    for it in items:
        nombre = it["nombre"][:20]
        linea1 = f"{nombre}\n".encode("cp437", "replace")
        linea2 = f"  {it['cantidad']} x {it['precio']:.2f} = {it['precio']*it['cantidad']:.2f}\n".encode("cp437", "replace")
        lineas.append(linea1)
        lineas.append(linea2)
    lineas.append(b"--------------------------------\n")
    lineas.append(f"Subtotal:        $ {subtotal:.2f}\n".encode("cp437", "replace"))
    lineas.append(f"Margen:          {margen:.1f} %\n".encode("cp437", "replace"))
    lineas.append(NEGRITA_ON + f"TOTAL:           $ {total_final:.2f}\n".encode("cp437", "replace") + NEGRITA_OFF)
    lineas.append(b"--------------------------------\n")
    
    msg = cfg.get("mensaje_ticket", "¡Gracias por su compra!")
    lineas.append(alineacion + f"{msg}\n\n\n".encode("cp437", "replace"))
    lineas.append(CORTE_PAPEL)
    return b"".join(lineas)


def imprimir_ticket(cfg, items, subtotal, margen, total_final):
    if not HAS_WIN32PRINT:
        messagebox.showwarning(
            "Impresora no disponible",
            "No se encontró el módulo 'pywin32'. Instalalo con:\n\n    pip install pywin32\n\n"
            "y volvé a intentar imprimir. Mientras tanto podés usar 'Guardar en Excel'.",
        )
        return False

    nombre_impresora = cfg.get("nombre_impresora", "").strip() or win32print.GetDefaultPrinter()
    data = construir_ticket(cfg, items, subtotal, margen, total_final)
    try:
        hprinter = win32print.OpenPrinter(nombre_impresora)
        try:
            hjob = win32print.StartDocPrinter(hprinter, 1, ("Ticket de venta", None, "RAW"))
            win32print.StartPagePrinter(hprinter)
            win32print.WritePrinter(hprinter, data)
            win32print.EndPagePrinter(hprinter)
            win32print.EndDocPrinter(hprinter)
        finally:
            win32print.ClosePrinter(hprinter)
        return True
    except Exception as e:
        messagebox.showerror("Error al imprimir", f"No se pudo imprimir el ticket:\n{e}")
        return False


def imprimir_normal_texto(texto_impresion, titulo_documento="Documento"):
    """Envía un texto plano a una impresora estándar (A4/Sistema Windows)."""
    if not HAS_WIN32PRINT:
        messagebox.showwarning(
            "Impresora no disponible",
            "No se encontró el módulo 'pywin32'. Instalalo con:\n\n    pip install pywin32",
        )
        return False

    try:
        hDC = win32ui.CreateDC()
        hDC.CreatePrinterDC(win32print.GetDefaultPrinter())
        hDC.StartDoc(titulo_documento)
        hDC.StartPage()

        font = win32ui.CreateFont({
            "name": "Arial",
            "height": 36,
            "weight": 400,
        })
        hDC.SelectObject(font)

        y = 100
        for linea in texto_impresion.split("\n"):
            hDC.TextOut(100, y, linea)
            y += 45

        hDC.EndPage()
        hDC.EndDoc()
        return True
    except Exception as e:
        messagebox.showerror("Error al imprimir", f"No se pudo imprimir en la impresora normal:\n{e}")
        return False


# ----------------------------------------------------------------------
# Excel: ventas.xlsx
# ----------------------------------------------------------------------
def guardar_venta_excel(items, margen=0.0):
    if openpyxl is None:
        messagebox.showwarning(
            "Excel no disponible",
            "No se encontró el módulo 'openpyxl'. Instalalo con:\n\n    pip install openpyxl",
        )
        return False

    if os.path.exists(VENTAS_XLSX):
        wb = load_workbook(VENTAS_XLSX)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"
        ws.append(VENTAS_HEADERS)

    ahora = datetime.now()
    dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    dia_semana = dias[ahora.weekday()]
    for it in items:
        subtotal = round(it["precio"] * it["cantidad"], 2)
        ganancia = round(subtotal * (margen / 100), 2)
        ws.append([
            dia_semana,
            ahora.strftime("%d/%m/%Y"),
            ahora.strftime("%H:%M:%S"),
            it["nombre"],
            it["precio"],
            it["cantidad"],
            subtotal,
            ganancia,
        ])
    wb.save(VENTAS_XLSX)
    return True


def obtener_ventas_del_dia():
    """Lee ventas.xlsx y retorna las ventas registradas con el día de hoy."""
    if not openpyxl or not os.path.exists(VENTAS_XLSX):
        return []

    ventas_hoy = []
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")

    try:
        wb = load_workbook(VENTAS_XLSX)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 7:
                continue
            dia, fecha, hora, producto, precio_u, cantidad, subtotal = row[:7]
            ganancia = row[7] if len(row) > 7 else None
            if str(fecha).strip() == fecha_hoy:
                ventas_hoy.append({
                    "dia": dia,
                    "fecha": fecha,
                    "hora": hora,
                    "producto": producto,
                    "precio_unitario": float(precio_u or 0),
                    "cantidad": int(cantidad or 0),
                    "subtotal": float(subtotal or 0),
                    "ganancia": float(ganancia) if ganancia is not None else None,
                })
    except Exception as e:
        print(f"Error al leer ventas de hoy: {e}")

    return ventas_hoy


# ----------------------------------------------------------------------
# Reportes de stock (xlsx + txt)
# ----------------------------------------------------------------------
def generar_reportes_stock():
    productos = listar_productos()
    disponibles = [p for p in productos if p[3] > 0]
    sin_stock = [p for p in productos if p[3] <= 0]

    if openpyxl is not None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Stock"
        ws.append(["Código", "Nombre", "Precio", "Stock"])
        for p in productos:
            ws.append(list(p))
        wb.save(STOCK_XLSX)

    with open(STOCK_TXT, "w", encoding="utf-8") as f:
        f.write(f"Reporte de stock - {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write("=" * 50 + "\n\n")
        f.write("PRODUCTOS DISPONIBLES\n")
        f.write("-" * 50 + "\n")
        for p in disponibles:
            f.write(f"{p[1]}  |  ${p[2]:.2f}  |  Stock: {p[3]}\n")
        f.write("\nPRODUCTOS SIN STOCK\n")
        f.write("-" * 50 + "\n")
        for p in sin_stock:
            f.write(f"{p[1]}  |  ${p[2]:.2f}\n")

    return STOCK_XLSX, STOCK_TXT


# ----------------------------------------------------------------------
# Paletas de tema (colores tomados del logo)
# ----------------------------------------------------------------------
LOGO_TEAL = "#4dbf9e"       # verde azulado del logo
LOGO_BLUE = "#007877"       # azul profundo del logo
LOGO_NAVY = "#0a296c"       # azul navy del logo
LOGO_DARK = "#122235"       # esquina oscura del logo (fondo dark)

TEMAS = {
    "claro": {
        "bg": "#eef5f4",            # fondo general, blanco con tinte verdoso muy suave
        "fg": "#12283a",            # texto principal, azul muy oscuro
        "panel": "#ffffff",         # tablas / paneles
        "accent": LOGO_BLUE,        # botones y acentos
        "accent_fg": "#ffffff",
        "accent2": LOGO_TEAL,       # segundo acento (para degradés y detalles)
        "borde": "#c9e2dd",
    },
    "oscuro": {
        "bg": LOGO_DARK,             # fondo general, navy oscuro del logo
        "fg": "#eaf5f2",             # texto claro
        "panel": "#1c3348",          # tablas / paneles
        "accent": LOGO_TEAL,         # botones y acentos (resalta sobre el fondo oscuro)
        "accent_fg": "#0a2018",
        "accent2": LOGO_BLUE,
        "borde": "#2a4a63",
    },
}

GRADIENTE_LOBBY = (LOGO_TEAL, LOGO_NAVY)


def _hex_a_rgb(hexcolor):
    hexcolor = hexcolor.lstrip("#")
    return tuple(int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))


def _interpolar_color(c1, c2, t):
    r1, g1, b1 = _hex_a_rgb(c1)
    r2, g2, b2 = _hex_a_rgb(c2)
    r = int(r1 + (r2 - r1) * t)
    g = int(g1 + (g2 - g1) * t)
    b = int(b1 + (b2 - b1) * t)
    return f"#{r:02x}{g:02x}{b:02x}"


class BotonDegradado(tk.Canvas):
    """Cuadrado con relleno en degradé (estilo del logo) que actúa como botón."""

    def __init__(self, parent, texto, comando, color_inicio, color_fin, fondo,
                 ancho=180, alto=180, fuente=("Segoe UI", 13, "bold")):
        super().__init__(parent, width=ancho, height=alto, highlightthickness=0, bg=fondo, bd=0)
        self.comando = comando
        self._dibujar_degradado(ancho, alto, color_inicio, color_fin)
        self.create_text(
            ancho / 2, alto / 2, text=texto, fill="#ffffff", font=fuente,
            width=ancho - 30, justify="center",
        )
        self.configure(cursor="hand2")
        self.bind("<Button-1>", lambda e: self.comando())

    def _dibujar_degradado(self, ancho, alto, color_inicio, color_fin):
        for i in range(ancho):
            t = i / max(1, (ancho - 1))
            color = _interpolar_color(color_inicio, color_fin, t)
            self.create_line(i, 0, i, alto, fill=color)
        self.create_rectangle(0, alto - 6, ancho, alto, fill=color_fin, outline=color_fin)


# ----------------------------------------------------------------------
# Aplicación principal
# ----------------------------------------------------------------------
class POSApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.cfg = cargar_config()
        self.title("Sistema de Punto de Venta")
        self.geometry("1000x680")
        self.minsize(880, 600)

        self.style = ttk.Style(self)
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass

        self.aplicar_tema()

        self.contenedor = tk.Frame(self, bg=self.paleta["bg"])
        self.contenedor.pack(fill="both", expand=True)

        self.frame_actual = None
        self.mostrar_lobby()

    def aplicar_tema(self):
        self.paleta = TEMAS.get(self.cfg.get("tema", "claro"), TEMAS["claro"])
        tam = self.cfg.get("tamano_fuente", 11)
        self.fuente_normal = ("Segoe UI", tam)
        self.fuente_titulo = ("Segoe UI", tam + 8, "bold")
        self.fuente_boton = ("Segoe UI", tam + 1, "bold")
        self.configure(bg=self.paleta["bg"])

        self.style.configure("TFrame", background=self.paleta["bg"])
        self.style.configure("Panel.TFrame", background=self.paleta["panel"])
        self.style.configure("TLabel", background=self.paleta["bg"], foreground=self.paleta["fg"], font=self.fuente_normal)
        self.style.configure("Panel.TLabel", background=self.paleta["panel"], foreground=self.paleta["fg"], font=self.fuente_normal)
        self.style.configure("Titulo.TLabel", background=self.paleta["bg"], foreground=self.paleta["fg"], font=self.fuente_titulo)
        self.style.configure("TButton", font=self.fuente_boton, padding=8)
        self.style.configure(
            "Treeview",
            background=self.paleta["panel"],
            fieldbackground=self.paleta["panel"],
            foreground=self.paleta["fg"],
            font=self.fuente_normal,
            rowheight=int(tam * 2.2),
        )
        self.style.configure(
            "Treeview.Heading", font=(self.fuente_normal[0], tam, "bold"),
            background=self.paleta["accent2"], foreground="#ffffff",
        )
        self.style.map("Treeview", background=[("selected", self.paleta["accent"])],
                        foreground=[("selected", self.paleta["accent_fg"])])
        self.style.configure("TEntry", font=self.fuente_normal)

    def refrescar_todo(self):
        self.aplicar_tema()
        self.contenedor.configure(bg=self.paleta["bg"])
        self.mostrar_lobby()

    def cambiar_frame(self, frame_cls):
        if self.frame_actual is not None:
            self.frame_actual.destroy()
        self.frame_actual = frame_cls(self.contenedor, self)
        self.frame_actual.pack(fill="both", expand=True)

    def mostrar_lobby(self):
        self.cambiar_frame(LobbyFrame)

    def mostrar_nueva_venta(self):
        self.cambiar_frame(NuevaVentaFrame)

    def mostrar_stock(self):
        self.cambiar_frame(StockFrame)

    def mostrar_cierre_caja(self):
        self.cambiar_frame(CierreCajaFrame)

    def mostrar_configuracion(self):
        self.cambiar_frame(ConfiguracionFrame)


# ----------------------------------------------------------------------
# Lobby
# ----------------------------------------------------------------------
class LobbyFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self._logo_ref = None

        if HAS_PIL and os.path.exists(LOGO_PATH):
            try:
                img = Image.open(LOGO_PATH).convert("RGBA").resize((96, 96), Image.LANCZOS)
                self._logo_ref = ImageTk.PhotoImage(img)
                tk.Label(self, image=self._logo_ref, bg=app.paleta["bg"]).pack(pady=(35, 5))
            except Exception:
                pass

        tk.Label(
            self, text="Punto de Venta", font=app.fuente_titulo,
            bg=app.paleta["bg"], fg=app.paleta["fg"]
        ).pack(pady=(5 if self._logo_ref else 50, 30))

        botones = tk.Frame(self, bg=app.paleta["bg"])
        botones.pack(expand=True)

        color_inicio, color_fin = GRADIENTE_LOBBY
        opciones = [
            ("Nueva\nVenta", app.mostrar_nueva_venta),
            ("Administrar\nStock", app.mostrar_stock),
            ("Cierre de\nCaja", app.mostrar_cierre_caja),
            ("Configuración", app.mostrar_configuracion),
        ]

        for i, (texto, comando) in enumerate(opciones):
            cuadro = BotonDegradado(
                botones, texto, comando, color_inicio, color_fin,
                fondo=app.paleta["bg"], fuente=(app.fuente_titulo[0], app.fuente_normal[1] + 3, "bold"),
            )
            cuadro.grid(row=0, column=i, padx=12, pady=20)


# ----------------------------------------------------------------------
# Cierre de Caja
# ----------------------------------------------------------------------
class CierreCajaFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self.ventas_hoy = obtener_ventas_del_dia()

        # Top bar
        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal).pack(side="left")
        tk.Label(top, text="Cierre de Caja del Día", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        # Tabla de ventas del día
        tabla_frame = tk.Frame(self, bg=app.paleta["bg"])
        tabla_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("hora", "producto", "precio", "cantidad", "subtotal")
        self.tree = ttk.Treeview(tabla_frame, columns=columnas, show="headings", selectmode="browse")
        for col, texto, ancho in [
            ("hora", "Hora", 90), ("producto", "Producto", 280),
            ("precio", "Precio Unit.", 110), ("cantidad", "Cant.", 80),
            ("subtotal", "Subtotal", 120)
        ]:
            self.tree.heading(col, text=texto)
            self.tree.column(col, width=ancho, anchor="center")

        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(tabla_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)

        # Totales
        totales_frame = tk.Frame(self, bg=app.paleta["bg"])
        totales_frame.pack(fill="x", padx=20, pady=10)

        self.total_ganado = sum(v["subtotal"] for v in self.ventas_hoy)
        
        # Ganancia estimada = suma de la ganancia calculada con el margen (%)
        # cargado en "Nueva venta" al finalizar cada compra.
        self.ganancia_estimada = 0.0
        for v in self.ventas_hoy:
            if v.get("ganancia") is not None:
                self.ganancia_estimada += v["ganancia"]
            else:
                # Ventas viejas guardadas antes de tener columna "Ganancia"
                prod = buscar_producto_por_nombre(v["producto"])
                if prod:
                    costo_base = prod[2]
                    subtotal_costo = costo_base * v["cantidad"]
                    self.ganancia_estimada += max(0.0, v["subtotal"] - subtotal_costo)
                else:
                    self.ganancia_estimada += v["subtotal"] * 0.20

        self.lbl_total_ganado = tk.Label(
            totales_frame, text=f"Total Ganado Hoy (Ventas): $ {self.total_ganado:.2f}",
            font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["accent"]
        )
        self.lbl_total_ganado.pack(anchor="w")

        self.lbl_ganancia_neta = tk.Label(
            totales_frame, text=f"Ganancia Estimada Hoy: $ {self.ganancia_estimada:.2f}",
            font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]
        )
        self.lbl_ganancia_neta.pack(anchor="w", pady=(5, 0))

        # Botones de Acción
        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))

        tk.Button(acciones_frame, text="💾 Guardar Cierre en Excel", command=self.guardar_excel, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🖨 Imprimir (Térmica)", command=self.imprimir_termica, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🖨 Imprimir (Normal)", command=self.imprimir_normal, font=app.fuente_boton,
                  bg=LOGO_TEAL, fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)

        self.cargar_tabla()

    def cargar_tabla(self):
        self.tree.delete(*self.tree.get_children())
        for v in self.ventas_hoy:
            self.tree.insert("", "end", values=(
                v["hora"], v["producto"], f"{v['precio_unitario']:.2f}", v["cantidad"], f"{v['subtotal']:.2f}"
            ))

    def guardar_excel(self):
        if not openpyxl:
            messagebox.showwarning("Excel no disponible", "No se encontró el módulo 'openpyxl'.")
            return
        if not self.ventas_hoy:
            messagebox.showinfo("Sin datos", "No hay ventas registradas el día de hoy para exportar.")
            return

        nombre_archivo = f"cierre_caja_{datetime.now().strftime('%Y%m%d')}.xlsx"
        path_cierre = os.path.join(BASE_DIR, nombre_archivo)

        wb = Workbook()
        ws = wb.active
        ws.title = "Cierre de Caja"
        ws.append(["Hora", "Producto", "Precio Unitario", "Cantidad", "Subtotal"])

        for v in self.ventas_hoy:
            ws.append([v["hora"], v["producto"], v["precio_unitario"], v["cantidad"], v["subtotal"]])

        ws.append([])
        ws.append(["TOTAL VENDIDO HOY:", self.total_ganado])
        ws.append(["GANANCIA ESTIMADA HOY:", self.ganancia_estimada])

        wb.save(path_cierre)
        messagebox.showinfo("Cierre Guardado", f"El informe de cierre de caja se guardó en:\n{path_cierre}")

    def imprimir_termica(self):
        if not self.ventas_hoy:
            messagebox.showinfo("Sin ventas", "No hay ventas registradas hoy para imprimir.")
            return

        items_cierre = [{
            "codigo": "",
            "nombre": v["producto"],
            "precio": v["precio_unitario"],
            "cantidad": v["cantidad"]
        } for v in self.ventas_hoy]

        imprimir_ticket(self.app.cfg, items_cierre, self.total_ganado, 0.0, self.total_ganado)

    def imprimir_normal(self):
        if not self.ventas_hoy:
            messagebox.showinfo("Sin ventas", "No hay ventas registradas hoy para imprimir.")
            return

        fecha_str = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        lineas = []
        lineas.append(f"=========================================")
        lineas.append(f"       REPORTE DE CIERRE DE CAJA        ")
        lineas.append(f" Local: {self.app.cfg.get('nombre_local', 'MI COMERCIO')}")
        lineas.append(f" Fecha: {fecha_str}")
        lineas.append(f"=========================================")
        lineas.append(f"{'Hora':<8} {'Producto':<20} {'Cant.':<6} {'Subtotal':<10}")
        lineas.append(f"-----------------------------------------")
        for v in self.ventas_hoy:
            lineas.append(f"{v['hora']:<8} {v['producto'][:18]:<20} {v['cantidad']:<6} ${v['subtotal']:.2f}")
        lineas.append(f"-----------------------------------------")
        lineas.append(f"TOTAL VENDIDO HOY:      $ {self.total_ganado:.2f}")
        lineas.append(f"GANANCIA ESTIMADA HOY:  $ {self.ganancia_estimada:.2f}")
        lineas.append(f"=========================================")
        lineas.append(f" {self.app.cfg.get('mensaje_ticket', '¡Gracias por su compra!')}")

        texto = "\n".join(lineas)
        if imprimir_normal_texto(texto, titulo_documento="Cierre_de_Caja"):
            messagebox.showinfo("Impresión", "Documento enviado a la impresora predeterminada.")


# ----------------------------------------------------------------------
# Nueva Venta
# ----------------------------------------------------------------------
class NuevaVentaFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self.items = {}

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal).pack(side="left")
        tk.Label(top, text="Nueva Venta", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        scan_frame = tk.Frame(self, bg=app.paleta["bg"])
        scan_frame.pack(fill="x", padx=20, pady=10)
        tk.Label(scan_frame, text="Escanear código de barras:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left")
        self.entry_codigo = tk.Entry(scan_frame, font=app.fuente_normal, width=30)
        self.entry_codigo.pack(side="left", padx=10)
        self.entry_codigo.bind("<Return>", self.escanear)
        self.entry_codigo.focus_set()
        tk.Button(scan_frame, text="Agregar", command=self.escanear, font=app.fuente_normal).pack(side="left")

        tabla_frame = tk.Frame(self, bg=app.paleta["bg"])
        tabla_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("codigo", "nombre", "precio", "cantidad", "subtotal")
        self.tree = ttk.Treeview(tabla_frame, columns=columnas, show="headings", selectmode="browse")
        for col, texto, ancho in [
            ("codigo", "Código", 120), ("nombre", "Producto", 260),
            ("precio", "Precio unit.", 110), ("cantidad", "Cantidad", 90),
            ("subtotal", "Subtotal", 110),
        ]:
            self.tree.heading(col, text=texto)
            self.tree.column(col, width=ancho, anchor="center")
        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(tabla_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.bind("<Double-1>", self.editar_cantidad)

        totales_frame = tk.Frame(self, bg=app.paleta["bg"])
        totales_frame.pack(fill="x", padx=20, pady=10)

        tk.Label(totales_frame, text="Margen de ganancia (%):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=0, column=0, sticky="w")
        self.entry_margen = tk.Entry(totales_frame, font=app.fuente_normal, width=8)
        self.entry_margen.insert(0, "0")
        self.entry_margen.grid(row=0, column=1, padx=10, sticky="w")
        self.entry_margen.bind("<KeyRelease>", lambda e: self.recalcular())

        self.lbl_subtotal = tk.Label(totales_frame, text="Subtotal: $ 0.00", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"])
        self.lbl_subtotal.grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))
        self.lbl_total = tk.Label(totales_frame, text="TOTAL: $ 0.00", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["accent"])
        self.lbl_total.grid(row=2, column=0, columnspan=2, sticky="w")

        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))

        tk.Button(acciones_frame, text="🖨 Imprimir ticket", command=self.imprimir, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="💾 Guardar en Excel", command=self.guardar_excel, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="✔ Finalizar compra", command=self.finalizar, font=app.fuente_boton,
                  bg=LOGO_TEAL, fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🗑 Eliminar producto", command=self.eliminar_producto, font=app.fuente_boton,
                  bg="#d94c4c", fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)

    def escanear(self, event=None):
        codigo = self.entry_codigo.get().strip()
        self.entry_codigo.delete(0, tk.END)
        if not codigo:
            return
        producto = buscar_producto(codigo)
        if producto is None:
            messagebox.showwarning("No encontrado", f"El código '{codigo}' no está registrado en Stock.")
            return
        _, nombre, precio, stock = producto
        if codigo in self.items:
            self.items[codigo]["cantidad"] += 1
        else:
            self.items[codigo] = {"codigo": codigo, "nombre": nombre, "precio": precio, "cantidad": 1}
        self.refrescar_tabla()

    def refrescar_tabla(self):
        self.tree.delete(*self.tree.get_children())
        for it in self.items.values():
            subtotal = it["precio"] * it["cantidad"]
            self.tree.insert("", "end", iid=it["codigo"], values=(
                it["codigo"], it["nombre"], f"{it['precio']:.2f}", it["cantidad"], f"{subtotal:.2f}",
            ))
        self.recalcular()

    def recalcular(self):
        subtotal = sum(it["precio"] * it["cantidad"] for it in self.items.values())
        try:
            margen = float(self.entry_margen.get().replace(",", "."))
        except ValueError:
            margen = 0.0
        total = subtotal * (1 + margen / 100)
        self.lbl_subtotal.config(text=f"Subtotal: $ {subtotal:.2f}")
        self.lbl_total.config(text=f"TOTAL: $ {total:.2f}")
        return subtotal, margen, total

    def editar_cantidad(self, event):
        item_id = self.tree.identify_row(event.y)
        columna = self.tree.identify_column(event.x)
        if not item_id or columna != "#4":
            return
        actual = self.items[item_id]["cantidad"]
        nueva = simpledialog.askinteger("Editar cantidad", f"Cantidad para '{self.items[item_id]['nombre']}':",
                                         initialvalue=actual, minvalue=1)
        if nueva:
            self.items[item_id]["cantidad"] = nueva
            self.refrescar_tabla()

    def eliminar_producto(self):
        seleccion = self.tree.selection()
        if not seleccion:
            messagebox.showinfo("Eliminar producto", "Seleccioná un producto de la lista primero.")
            return
        codigo = seleccion[0]
        item = self.items[codigo]
        if item["cantidad"] > 1:
            cantidad_a_borrar = simpledialog.askinteger(
                "Eliminar producto",
                f"'{item['nombre']}' tiene {item['cantidad']} unidades cargadas.\n"
                f"¿Cuántas unidades querés quitar?",
                initialvalue=item["cantidad"], minvalue=1, maxvalue=item["cantidad"],
            )
            if not cantidad_a_borrar:
                return
            if cantidad_a_borrar >= item["cantidad"]:
                del self.items[codigo]
            else:
                item["cantidad"] -= cantidad_a_borrar
        else:
            del self.items[codigo]
        self.refrescar_tabla()

    def _items_como_lista(self):
        return list(self.items.values())

    def imprimir(self):
        if not self.items:
            messagebox.showinfo("Ticket vacío", "Todavía no escaneaste ningún producto.")
            return
        subtotal, margen, total = self.recalcular()
        imprimir_ticket(self.app.cfg, self._items_como_lista(), subtotal, margen, total)

    def guardar_excel(self):
        if not self.items:
            messagebox.showinfo("Ticket vacío", "Todavía no escaneaste ningún producto.")
            return
        _, margen, _ = self.recalcular()
        if guardar_venta_excel(self._items_como_lista(), margen):
            messagebox.showinfo("Guardado", f"La venta se guardó en:\n{VENTAS_XLSX}")

    def finalizar(self):
        if not self.items:
            messagebox.showinfo("Venta vacía", "Todavía no escaneaste ningún producto.")
            return
        _, margen, _ = self.recalcular()
        guardar_venta_excel(self._items_como_lista(), margen)
        umbral = self.app.cfg.get("umbral_stock_bajo", 3)
        for it in self.items.values():
            nombre, restante = descontar_stock(it["codigo"], it["cantidad"])
            if nombre is not None and restante <= umbral:
                enviar_alerta_stock(self.app.cfg, nombre, restante)
        messagebox.showinfo("Compra finalizada", "La venta se registró y el stock fue actualizado.")
        self.app.mostrar_lobby()


# ----------------------------------------------------------------------
# Administrar Stock
# ----------------------------------------------------------------------
class StockFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal).pack(side="left")
        tk.Label(top, text="Administrar Stock", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        alta_frame = tk.LabelFrame(self, text="Registrar / actualizar producto", font=app.fuente_normal,
                                    bg=app.paleta["bg"], fg=app.paleta["fg"])
        alta_frame.pack(fill="x", padx=20, pady=10)

        tk.Label(alta_frame, text="Código de barras:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=0, column=0, padx=5, pady=8, sticky="e")
        self.entry_codigo = tk.Entry(alta_frame, font=app.fuente_normal, width=20)
        self.entry_codigo.grid(row=0, column=1, padx=5)
        self.entry_codigo.bind("<Return>", self.buscar_para_editar)

        tk.Label(alta_frame, text="Nombre:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=0, column=2, padx=5, sticky="e")
        self.entry_nombre = tk.Entry(alta_frame, font=app.fuente_normal, width=25)
        self.entry_nombre.grid(row=0, column=3, padx=5)

        tk.Label(alta_frame, text="Precio:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=1, column=0, padx=5, sticky="e")
        self.entry_precio = tk.Entry(alta_frame, font=app.fuente_normal, width=20)
        self.entry_precio.grid(row=1, column=1, padx=5)

        self.var_label_stock = tk.StringVar(value="Stock inicial:")
        tk.Label(alta_frame, textvariable=self.var_label_stock, bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=1, column=2, padx=5, sticky="e")
        self.entry_stock = tk.Entry(alta_frame, font=app.fuente_normal, width=25)
        self.entry_stock.grid(row=1, column=3, padx=5)
        self.entry_stock.bind("<Return>", lambda e: self.guardar())

        tk.Button(alta_frame, text="Guardar producto", command=self.guardar, font=app.fuente_normal).grid(row=0, column=4, rowspan=2, padx=15)

        self.lbl_info = tk.Label(alta_frame, text="", font=(app.fuente_normal[0], app.fuente_normal[1] - 1),
                                  bg=app.paleta["bg"], fg=app.paleta["accent"], justify="left")
        self.lbl_info.grid(row=2, column=0, columnspan=5, sticky="w", padx=5, pady=(0, 8))

        lista_frame = tk.Frame(self, bg=app.paleta["bg"])
        lista_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("nombre", "precio", "stock", "codigo")
        self.tree = ttk.Treeview(lista_frame, columns=columnas, show="headings", selectmode="browse")
        self.tree.heading("nombre", text="Producto")
        self.tree.heading("precio", text="Precio")
        self.tree.heading("stock", text="Stock")
        self.tree.heading("codigo", text="Código")
        self.tree.column("nombre", width=320)
        self.tree.column("precio", width=100, anchor="center")
        self.tree.column("stock", width=100, anchor="center")
        self.tree.column("codigo", width=140, anchor="center")
        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(lista_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.bind("<<TreeviewSelect>>", self.cargar_seleccion)

        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))
        tk.Button(acciones_frame, text="Eliminar producto seleccionado", command=self.eliminar, font=app.fuente_normal).pack(side="left", padx=5)
        tk.Button(acciones_frame, text="📄 Generar reporte de stock (Excel + TXT)", command=self.generar_reporte, font=app.fuente_normal).pack(side="left", padx=5)

        self.cargar_lista()

    def cargar_lista(self):
        self.tree.delete(*self.tree.get_children())
        for codigo, nombre, precio, stock in listar_productos():
            self.tree.insert("", "end", iid=codigo, values=(nombre, f"{precio:.2f}", stock, codigo))

    def buscar_para_editar(self, event=None):
        codigo = self.entry_codigo.get().strip()
        if not codigo:
            return
        self._cargar_en_formulario(codigo)
        self.entry_stock.focus_set()

    def cargar_seleccion(self, event=None):
        seleccion = self.tree.selection()
        if not seleccion:
            return
        codigo = seleccion[0]
        self.entry_codigo.delete(0, tk.END)
        self.entry_codigo.insert(0, codigo)
        self._cargar_en_formulario(codigo)

    def _cargar_en_formulario(self, codigo):
        producto = buscar_producto(codigo)
        self.entry_nombre.delete(0, tk.END)
        self.entry_precio.delete(0, tk.END)
        self.entry_stock.delete(0, tk.END)
        if producto:
            _, nombre, precio, stock = producto
            self.entry_nombre.insert(0, nombre)
            self.entry_precio.insert(0, str(precio))
            self.var_label_stock.set("Cantidad a agregar:")
            self.lbl_info.config(
                text=f"✔ Producto ya registrado — stock actual: {stock} unidades. "
                     f"La cantidad que cargues acá se SUMARÁ a ese stock (dejalo vacío para no tocar el stock).",
            )
        else:
            self.var_label_stock.set("Stock inicial:")
            self.lbl_info.config(text="🆕 Producto nuevo — completá nombre, precio y stock inicial.")
            self.entry_nombre.focus_set()

    def guardar(self):
        codigo = self.entry_codigo.get().strip()
        nombre = self.entry_nombre.get().strip()
        precio_txt = self.entry_precio.get().strip().replace(",", ".")
        stock_txt = self.entry_stock.get().strip()
        if not codigo or not nombre or not precio_txt:
            messagebox.showwarning("Datos incompletos", "Completá código, nombre y precio.")
            return
        try:
            precio = float(precio_txt)
        except ValueError:
            messagebox.showwarning("Datos inválidos", "El precio debe ser numérico.")
            return

        cantidad_ingresada = 0
        if stock_txt:
            try:
                cantidad_ingresada = int(stock_txt)
            except ValueError:
                messagebox.showwarning("Datos inválidos", "La cantidad/stock debe ser un número entero.")
                return

        producto_existente = buscar_producto(codigo)
        if producto_existente:
            stock_actual = producto_existente[3]
            nuevo_stock = stock_actual + cantidad_ingresada
            guardar_producto(codigo, nombre, precio, nuevo_stock)
            if cantidad_ingresada:
                messagebox.showinfo(
                    "Stock actualizado",
                    f"'{nombre}': {stock_actual} + {cantidad_ingresada} = {nuevo_stock} unidades.",
                )
        else:
            if not stock_txt:
                messagebox.showwarning("Datos incompletos", "Para un producto nuevo indicá el stock inicial.")
                return
            guardar_producto(codigo, nombre, precio, cantidad_ingresada)

        self.cargar_lista()
        self.entry_codigo.delete(0, tk.END)
        self.entry_nombre.delete(0, tk.END)
        self.entry_precio.delete(0, tk.END)
        self.entry_stock.delete(0, tk.END)
        self.var_label_stock.set("Stock inicial:")
        self.lbl_info.config(text="")
        self.entry_codigo.focus_set()

    def eliminar(self):
        seleccion = self.tree.selection()
        if not seleccion:
            messagebox.showinfo("Eliminar producto", "Seleccioná un producto de la lista primero.")
            return
        codigo = seleccion[0]
        if messagebox.askyesno("Confirmar", "¿Eliminar este producto del stock?"):
            eliminar_producto(codigo)
            self.cargar_lista()

    def generar_reporte(self):
        xlsx_path, txt_path = generar_reportes_stock()
        messagebox.showinfo("Reporte generado", f"Se generaron:\n{xlsx_path}\n{txt_path}")


# ----------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------
class ConfiguracionFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal).pack(side="left")
        tk.Label(top, text="Configuración", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        panel = tk.Frame(self, bg=app.paleta["bg"])
        panel.pack(fill="both", expand=True, padx=30, pady=10)

        # Tema
        tk.Label(panel, text="Tema:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=0, column=0, sticky="w", pady=5)
        self.var_tema = tk.StringVar(value=app.cfg.get("tema", "claro"))
        tk.Radiobutton(panel, text="Claro", variable=self.var_tema, value="claro", font=app.fuente_normal,
                        bg=app.paleta["bg"], fg=app.paleta["fg"], selectcolor=app.paleta["panel"]).grid(row=0, column=1, sticky="w")
        tk.Radiobutton(panel, text="Oscuro", variable=self.var_tema, value="oscuro", font=app.fuente_normal,
                        bg=app.paleta["bg"], fg=app.paleta["fg"], selectcolor=app.paleta["panel"]).grid(row=0, column=2, sticky="w")

        # Tamaño de letra
        tk.Label(panel, text="Tamaño de letra:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=1, column=0, sticky="w", pady=5)
        self.var_tamano = tk.IntVar(value=app.cfg.get("tamano_fuente", 11))
        tk.Spinbox(panel, from_=8, to=20, textvariable=self.var_tamano, width=5, font=app.fuente_normal).grid(row=1, column=1, sticky="w")

        # Umbral de stock bajo
        tk.Label(panel, text="Umbral de stock bajo (unidades):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=2, column=0, sticky="w", pady=5)
        self.var_umbral = tk.IntVar(value=app.cfg.get("umbral_stock_bajo", 3))
        tk.Spinbox(panel, from_=1, to=100, textvariable=self.var_umbral, width=5, font=app.fuente_normal).grid(row=2, column=1, sticky="w")

        # --- Edición de Texto de Impresión (Petición) ---
        tk.Label(panel, text="Personalización de Impresión", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=3, column=0, columnspan=3, sticky="w", pady=(15, 5))

        tk.Label(panel, text="Nombre del Local:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=4, column=0, sticky="w", pady=5)
        self.entry_nombre_local = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_nombre_local.insert(0, app.cfg.get("nombre_local", "MI COMERCIO"))
        self.entry_nombre_local.grid(row=4, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Mensaje en Ticket:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=5, column=0, sticky="w", pady=5)
        self.entry_mensaje_ticket = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_mensaje_ticket.insert(0, app.cfg.get("mensaje_ticket", "¡Gracias por su compra!"))
        self.entry_mensaje_ticket.grid(row=5, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Alineación del Encabezado:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=6, column=0, sticky="w", pady=5)
        self.var_disposicion = tk.StringVar(value=app.cfg.get("disposicion_ticket", "Centro"))
        combo_disp = ttk.Combobox(panel, textvariable=self.var_disposicion, values=["Izquierda", "Centro", "Derecha"], state="readonly", width=15)
        combo_disp.grid(row=6, column=1, sticky="w")

        # --- Alertas por email ---
        tk.Label(panel, text="Avisos de stock bajo por email", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=7, column=0, columnspan=3, sticky="w", pady=(15, 5))

        tk.Label(panel, text="Email que recibe el aviso:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=8, column=0, sticky="w", pady=5)
        self.entry_email_alertas = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_email_alertas.insert(0, app.cfg.get("email_alertas", ""))
        self.entry_email_alertas.grid(row=8, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Cuenta Gmail que envía (remitente):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=9, column=0, sticky="w", pady=5)
        self.entry_email_remitente = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_email_remitente.insert(0, app.cfg.get("email_remitente", ""))
        self.entry_email_remitente.grid(row=9, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Contraseña de aplicación de Gmail:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=10, column=0, sticky="w", pady=5)
        self.entry_email_password = tk.Entry(panel, font=app.fuente_normal, width=35, show="*")
        self.entry_email_password.insert(0, app.cfg.get("email_password", ""))
        self.entry_email_password.grid(row=10, column=1, columnspan=2, sticky="w")

        # --- Impresora ---
        tk.Label(panel, text="Nombre exacto de la impresora (opcional):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=11, column=0, sticky="w", pady=5)
        self.entry_impresora = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_impresora.insert(0, app.cfg.get("nombre_impresora", ""))
        self.entry_impresora.grid(row=11, column=1, columnspan=2, sticky="w")

        # --- Guardar ---
        tk.Button(panel, text="Guardar configuración", font=app.fuente_boton, command=self.guardar,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").grid(row=12, column=0, pady=15, ipadx=10, ipady=6, sticky="w")

    def guardar(self):
        self.app.cfg["tema"] = self.var_tema.get()
        self.app.cfg["tamano_fuente"] = int(self.var_tamano.get())
        self.app.cfg["umbral_stock_bajo"] = int(self.var_umbral.get())
        self.app.cfg["nombre_local"] = self.entry_nombre_local.get().strip()
        self.app.cfg["mensaje_ticket"] = self.entry_mensaje_ticket.get().strip()
        self.app.cfg["disposicion_ticket"] = self.var_disposicion.get()
        self.app.cfg["email_alertas"] = self.entry_email_alertas.get().strip()
        self.app.cfg["email_remitente"] = self.entry_email_remitente.get().strip()
        self.app.cfg["email_password"] = self.entry_email_password.get().strip()
        self.app.cfg["nombre_impresora"] = self.entry_impresora.get().strip()
        guardar_config(self.app.cfg)
        messagebox.showinfo("Configuración", "Los cambios se guardaron correctamente.")
        self.app.refrescar_todo()


# ----------------------------------------------------------------------
# Punto de entrada
# ----------------------------------------------------------------------
def main():
    app = POSApp()
    app.mainloop()


if __name__ == "__main__":
    main()