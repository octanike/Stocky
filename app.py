import json
import os
import sqlite3
import smtplib
import threading
import sys
import time
import urllib.request
import urllib.error
from datetime import datetime
from email.mime.text import MIMEText

import tkinter as tk
from tkinter import ttk, messagebox, simpledialog

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
except ImportError:
    openpyxl = None

try:
    import win32print
    import win32ui
    import win32con
    HAS_WIN32PRINT = True
except ImportError:
    HAS_WIN32PRINT = False

try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import matplotlib.pyplot as plt
except ImportError:
    plt = None


# ----------------------------------------------------------------------
# Rutas y constantes
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "stock.db")
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
VENTAS_XLSX = os.path.join(BASE_DIR, "ventas.xlsx")
STOCK_XLSX = os.path.join(BASE_DIR, "stock.xlsx")
STOCK_TXT = os.path.join(BASE_DIR, "stock.txt")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo.jpg")

DEFAULT_CONFIG = {
    "tema": "claro",              # "claro" u "oscuro"
    "tamano_fuente": 11,
    "email_alertas": "",          # a dónde se manda el aviso de stock bajo
    "email_remitente": "",        # cuenta de Gmail que envía el aviso
    "email_password": "",         # contraseña de aplicación de Gmail
    "smtp_server": "smtp.gmail.com",
    "smtp_port": 465,
    "umbral_stock_bajo": 3,
    "nombre_impresora": "",       # vacío = usar la impresora predeterminada
    "nombre_local": "MI COMERCIO",
    "mensaje_ticket": "¡Gracias por su compra!",
    "disposicion_ticket": "Centro", # "Izquierda", "Centro", "Derecha"
    "gemini_api_key": "",          # clave gratuita de Google AI Studio (aistudio.google.com/apikey)
    "gemini_model": "gemini-2.5-flash",  # modelo gratuito de Gemini
    "notificaciones_activas": True,           # mostrar u ocultar los avisos de stock bajo en pantalla
    "notificaciones_intervalo_minutos": 1.5,  # cada cuántos minutos se revisa el stock y aparecen los avisos
}

VENTAS_HEADERS = ["Día", "Fecha", "Hora", "Producto", "Precio Unitario", "Cantidad", "Subtotal", "Ganancia"]


# ----------------------------------------------------------------------
# Configuración (persistida en config.json)
# ----------------------------------------------------------------------
def cargar_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            cfg = DEFAULT_CONFIG.copy()
            cfg.update(data)
            return cfg
        except Exception:
            pass
    return DEFAULT_CONFIG.copy()


def guardar_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)


# ----------------------------------------------------------------------
# Base de datos (SQLite) - productos / stock
# ----------------------------------------------------------------------
def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS productos (
            codigo TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            precio REAL NOT NULL,
            stock INTEGER NOT NULL DEFAULT 0
        )
        """
    )
    return conn


def buscar_producto(codigo):
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos WHERE codigo = ?", (codigo,))
    row = cur.fetchone()
    conn.close()
    return row  # (codigo, nombre, precio, stock) o None


def buscar_producto_por_nombre(nombre):
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos WHERE nombre = ?", (nombre,))
    row = cur.fetchone()
    conn.close()
    return row


def guardar_producto(codigo, nombre, precio, stock):
    conn = get_conn()
    conn.execute(
        "INSERT INTO productos (codigo, nombre, precio, stock) VALUES (?,?,?,?) "
        "ON CONFLICT(codigo) DO UPDATE SET nombre=excluded.nombre, precio=excluded.precio, stock=excluded.stock",
        (codigo, nombre, precio, stock),
    )
    conn.commit()
    conn.close()


def eliminar_producto(codigo):
    conn = get_conn()
    conn.execute("DELETE FROM productos WHERE codigo = ?", (codigo,))
    conn.commit()
    conn.close()


def listar_productos():
    conn = get_conn()
    cur = conn.execute("SELECT codigo, nombre, precio, stock FROM productos ORDER BY nombre ASC")
    rows = cur.fetchall()
    conn.close()
    return rows


def descontar_stock(codigo, cantidad):
    """Descuenta stock y devuelve (nombre, stock_restante)."""
    conn = get_conn()
    cur = conn.execute("SELECT nombre, stock FROM productos WHERE codigo = ?", (codigo,))
    row = cur.fetchone()
    if row is None:
        conn.close()
        return None, None
    nombre, stock_actual = row
    nuevo_stock = max(0, stock_actual - cantidad)
    conn.execute("UPDATE productos SET stock = ? WHERE codigo = ?", (nuevo_stock, codigo))
    conn.commit()
    conn.close()
    return nombre, nuevo_stock


# ----------------------------------------------------------------------
# Email de alerta de stock bajo
# ----------------------------------------------------------------------
def enviar_alerta_stock(cfg, nombre_producto, stock_restante):
    """Envía el email en un hilo aparte para no trabar la interfaz."""
    destino = cfg.get("email_alertas", "").strip()
    remitente = cfg.get("email_remitente", "").strip()
    password = cfg.get("email_password", "").strip()
    if not destino or not remitente or not password:
        return  # alertas no configuradas, se ignora silenciosamente

    def _enviar():
        try:
            asunto = f"Stock bajo: {nombre_producto}"
            cuerpo = (
                f"El producto '{nombre_producto}' quedó con {stock_restante} unidad(es) "
                f"en stock (umbral configurado: {cfg.get('umbral_stock_bajo')})."
            )
            msg = MIMEText(cuerpo, "plain", "utf-8")
            msg["Subject"] = asunto
            msg["From"] = remitente
            msg["To"] = destino
            with smtplib.SMTP_SSL(cfg.get("smtp_server", "smtp.gmail.com"), int(cfg.get("smtp_port", 465))) as server:
                server.login(remitente, password)
                server.sendmail(remitente, [destino], msg.as_string())
        except Exception as e:
            print(f"[Aviso] No se pudo enviar el email de alerta: {e}")

    threading.Thread(target=_enviar, daemon=True).start()


# ----------------------------------------------------------------------
# Impresión térmica y normal
# ----------------------------------------------------------------------
ESC = b"\x1b"
GS = b"\x1d"
CORTE_PAPEL = GS + b"V" + b"\x00"      # corta el papel
INIT_PRINTER = ESC + b"@"              # reset impresora
CENTRAR = ESC + b"a" + b"\x01"
ALINEAR_IZQ = ESC + b"a" + b"\x00"
ALINEAR_DER = ESC + b"a" + b"\x02"
NEGRITA_ON = ESC + b"E" + b"\x01"
NEGRITA_OFF = ESC + b"E" + b"\x00"


def obtener_comando_alineacion(disposicion):
    if disposicion == "Izquierda":
        return ALINEAR_IZQ
    elif disposicion == "Derecha":
        return ALINEAR_DER
    return CENTRAR


def construir_ticket(cfg, items, subtotal, margen, total_final):
    """items: lista de dicts {codigo, nombre, precio, cantidad}"""
    lineas = []
    lineas.append(INIT_PRINTER)
    
    alineacion = obtener_comando_alineacion(cfg.get("disposicion_ticket", "Centro"))
    lineas.append(alineacion)
    
    nombre_local = cfg.get("nombre_local", "MI COMERCIO").upper()
    lineas.append(NEGRITA_ON + f"{nombre_local}\n".encode("cp437", "replace") + NEGRITA_OFF)
    lineas.append(b"TICKET DE VENTA\n")
    lineas.append(ALINEAR_IZQ)
    ahora = datetime.now()
    lineas.append(f"{ahora.strftime('%d/%m/%Y')}  {ahora.strftime('%H:%M:%S')}\n".encode("cp437", "replace"))
    lineas.append(b"--------------------------------\n")
    for it in items:
        nombre = it["nombre"][:20]
        linea1 = f"{nombre}\n".encode("cp437", "replace")
        linea2 = f"  {it['cantidad']} x {it['precio']:.2f} = {it['precio']*it['cantidad']:.2f}\n".encode("cp437", "replace")
        lineas.append(linea1)
        lineas.append(linea2)
    lineas.append(b"--------------------------------\n")
    lineas.append(f"Subtotal:        $ {subtotal:.2f}\n".encode("cp437", "replace"))
    lineas.append(f"Margen:          {margen:.1f} %\n".encode("cp437", "replace"))
    lineas.append(NEGRITA_ON + f"TOTAL:           $ {total_final:.2f}\n".encode("cp437", "replace") + NEGRITA_OFF)
    lineas.append(b"--------------------------------\n")
    
    msg = cfg.get("mensaje_ticket", "¡Gracias por su compra!")
    lineas.append(alineacion + f"{msg}\n\n\n".encode("cp437", "replace"))
    lineas.append(CORTE_PAPEL)
    return b"".join(lineas)


def imprimir_ticket(cfg, items, subtotal, margen, total_final):
    if not HAS_WIN32PRINT:
        messagebox.showwarning(
            "Impresora no disponible",
            "No se encontró el módulo 'pywin32'. Instalalo con:\n\n    pip install pywin32\n\n"
            "y volvé a intentar imprimir. Mientras tanto podés usar 'Guardar en Excel'.",
        )
        return False

    nombre_impresora = cfg.get("nombre_impresora", "").strip() or win32print.GetDefaultPrinter()
    data = construir_ticket(cfg, items, subtotal, margen, total_final)
    try:
        hprinter = win32print.OpenPrinter(nombre_impresora)
        try:
            hjob = win32print.StartDocPrinter(hprinter, 1, ("Ticket de venta", None, "RAW"))
            win32print.StartPagePrinter(hprinter)
            win32print.WritePrinter(hprinter, data)
            win32print.EndPagePrinter(hprinter)
            win32print.EndDocPrinter(hprinter)
        finally:
            win32print.ClosePrinter(hprinter)
        return True
    except Exception as e:
        messagebox.showerror("Error al imprimir", f"No se pudo imprimir el ticket:\n{e}")
        return False


def imprimir_normal_texto(texto_impresion, titulo_documento="Documento"):
    """Envía un texto plano a una impresora estándar (A4/Sistema Windows)."""
    if not HAS_WIN32PRINT:
        messagebox.showwarning(
            "Impresora no disponible",
            "No se encontró el módulo 'pywin32'. Instalalo con:\n\n    pip install pywin32",
        )
        return False

    try:
        hDC = win32ui.CreateDC()
        hDC.CreatePrinterDC(win32print.GetDefaultPrinter())
        hDC.StartDoc(titulo_documento)
        hDC.StartPage()

        font = win32ui.CreateFont({
            "name": "Arial",
            "height": 36,
            "weight": 400,
        })
        hDC.SelectObject(font)

        y = 100
        for linea in texto_impresion.split("\n"):
            hDC.TextOut(100, y, linea)
            y += 45

        hDC.EndPage()
        hDC.EndDoc()
        return True
    except Exception as e:
        messagebox.showerror("Error al imprimir", f"No se pudo imprimir en la impresora normal:\n{e}")
        return False


# ----------------------------------------------------------------------
# Excel: ventas.xlsx
# ----------------------------------------------------------------------
def guardar_venta_excel(items, margen=0.0):
    if openpyxl is None:
        messagebox.showwarning(
            "Excel no disponible",
            "No se encontró el módulo 'openpyxl'. Instalalo con:\n\n    pip install openpyxl",
        )
        return False

    if os.path.exists(VENTAS_XLSX):
        wb = load_workbook(VENTAS_XLSX)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"
        ws.append(VENTAS_HEADERS)

    ahora = datetime.now()
    dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    dia_semana = dias[ahora.weekday()]
    for it in items:
        subtotal = round(it["precio"] * it["cantidad"], 2)
        ganancia = round(subtotal * (margen / 100), 2)
        ws.append([
            dia_semana,
            ahora.strftime("%d/%m/%Y"),
            ahora.strftime("%H:%M:%S"),
            it["nombre"],
            it["precio"],
            it["cantidad"],
            subtotal,
            ganancia,
        ])
    wb.save(VENTAS_XLSX)
    return True


def obtener_ventas_del_dia():
    """Lee ventas.xlsx y retorna las ventas registradas con el día de hoy."""
    if not openpyxl or not os.path.exists(VENTAS_XLSX):
        return []

    ventas_hoy = []
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")

    try:
        wb = load_workbook(VENTAS_XLSX)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 7:
                continue
            dia, fecha, hora, producto, precio_u, cantidad, subtotal = row[:7]
            ganancia = row[7] if len(row) > 7 else None
            if str(fecha).strip() == fecha_hoy:
                ventas_hoy.append({
                    "dia": dia,
                    "fecha": fecha,
                    "hora": hora,
                    "producto": producto,
                    "precio_unitario": float(precio_u or 0),
                    "cantidad": int(cantidad or 0),
                    "subtotal": float(subtotal or 0),
                    "ganancia": float(ganancia) if ganancia is not None else None,
                })
    except Exception as e:
        print(f"Error al leer ventas de hoy: {e}")

    return ventas_hoy


def _texto_contiene(texto, *palabras):
    texto = (texto or "").lower()
    return all(palabra.lower() in texto for palabra in palabras)


def obtener_ventas_generales():
    if not openpyxl or not os.path.exists(VENTAS_XLSX):
        return []

    ventas = []
    try:
        wb = load_workbook(VENTAS_XLSX)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 7:
                continue
            dia, fecha, hora, producto, precio_u, cantidad, subtotal = row[:7]
            ganancia = row[7] if len(row) > 7 else None
            ventas.append({
                "dia": dia,
                "fecha": str(fecha).strip() if fecha is not None else "",
                "hora": hora,
                "producto": producto,
                "precio_unitario": float(precio_u or 0),
                "cantidad": int(cantidad or 0),
                "subtotal": float(subtotal or 0),
                "ganancia": float(ganancia) if ganancia is not None else None,
            })
    except Exception as e:
        print(f"Error al leer ventas generales: {e}")

    return ventas


def calcular_top_productos(ventas, top_n=5):
    contador = {}
    for venta in ventas:
        nombre = str(venta.get("producto") or "").strip()
        if not nombre:
            continue
        datos = contador.setdefault(nombre, {"cantidad": 0, "subtotal": 0.0})
        datos["cantidad"] += int(venta.get("cantidad", 0))
        datos["subtotal"] += float(venta.get("subtotal", 0.0))
    ordenados = sorted(contador.items(), key=lambda item: item[1]["cantidad"], reverse=True)
    return ordenados[:top_n]


def generar_recomendaciones(ventas, productos):
    if not ventas:
        return (
            "Basado en los datos actuales, aún no hay ventas registradas. "
            "Empieza por cargar productos con stock suficiente, promocionar los más populares y revisar precios competitivos. "
            "También es útil usar la función de alertas para no quedarte sin stock en tus productos más vendidos."
        )

    top_vendidos = calcular_top_productos(ventas, top_n=3)
    productos_vendidos = {venta["producto"] for venta in ventas if venta.get("producto")}
    stock_bajo = [p for p in productos if p[3] <= 3]
    sin_venta = [p for p in productos if p[1] not in productos_vendidos]

    partes = [
        "Desde mi punto de vista, el mejor enfoque es priorizar a tus productos más vendidos y mantenerlos siempre disponibles.",
    ]
    if top_vendidos:
        nombres_top = ", ".join([p[0] for p in top_vendidos])
        partes.append(f"Los productos que más se venden son: {nombres_top}. Asegúrate de tener buen stock y promoverlos con ofertas o paquetes.")
    if stock_bajo:
        nombres_stock_bajo = ", ".join([p[1] for p in stock_bajo[:4]])
        partes.append(f"Hay productos con stock bajo como: {nombres_stock_bajo}. Recarga antes de que se agoten.")
    if sin_venta:
        nombres_sin_venta = ", ".join([p[1] for p in sin_venta[:4]])
        partes.append(
            f"Productos con stock sin ventas recientes: {nombres_sin_venta}. Considera promociones o cambios de ubicación para mover ese stock."
        )
    partes.append(
        "También puedes usar los reportes y las gráficas para evaluar qué productos debes potenciar y cuáles ajustar en precio."
    )
    return " ".join(partes)


CONOCIMIENTO_APP = (
    "Esta aplicación administra productos en un inventario SQLite, guarda ventas en ventas.xlsx, "
    "genera reportes de stock y ventas, imprime tickets en impresoras térmicas y normales, "
    "envía alertas de stock bajo por email y permite configurar tema, tamaño de fuente, nombre del local y mensaje de ticket."
)


def obtener_resumen_app(ventas, productos):
    total_prod = len(productos)
    total_ventas = sum(v["subtotal"] for v in ventas)
    total_unidades = sum(v["cantidad"] for v in ventas)
    top = calcular_top_productos(ventas, top_n=3)
    top_text = ", ".join([f"{p[0]} ({p[1]['cantidad']} uds)" for p in top]) if top else "sin datos aún"
    return (
        f"La app contiene {total_prod} productos, ha registrado {len(ventas)} líneas de venta por ${total_ventas:.2f} y {total_unidades} unidades vendidas. "
        f"Top 3 actuales: {top_text}."
    )


GEMINI_URL_TMPL = "https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"


def construir_contexto_app():
    """Arma el contexto completo (funciones + datos actuales) que se le manda a Gemini
    para que el asistente sepa absolutamente todo sobre la app y sus datos en tiempo real."""
    ventas_totales = obtener_ventas_generales()
    productos = listar_productos()
    resumen_app = obtener_resumen_app(ventas_totales, productos)
    top_vendidos = calcular_top_productos(ventas_totales, top_n=5)
    top_texto = (
        ", ".join([f"{p[0]} ({p[1]['cantidad']} uds, ${p[1]['subtotal']:.2f})" for p in top_vendidos])
        if top_vendidos else "sin datos aún"
    )
    lineas_stock = "\n".join(
        [f"- {p[1]} (código {p[0]}): ${p[2]:.2f}, stock {p[3]} unidades" for p in productos]
    ) or "No hay productos cargados todavía."

    contexto = (
        "Sos el asistente de inteligencia artificial integrado en una aplicación de Punto de Venta (POS) "
        "de escritorio hecha en Python/Tkinter. Conocés absolutamente todo sobre esta aplicación: sus pantallas, "
        "sus funciones y los datos reales que se listan abajo. Respondé SIEMPRE en español, de forma clara, breve, "
        "concreta y amable, usando esos datos reales cuando la pregunta lo requiera. Si te preguntan algo que no "
        "tiene que ver con la app, respondé igual con normalidad pero recordando que tu rol principal es ayudar con esta app.\n\n"
        f"DESCRIPCIÓN GENERAL:\n{CONOCIMIENTO_APP}\n\n"
        "FUNCIONES DE LA APP:\n"
        "- Nueva Venta: registra ventas, descuenta stock automáticamente, calcula la ganancia según el margen "
        "indicado e imprime el ticket (impresora térmica o normal).\n"
        "- Administrar Stock: alta, baja y edición de productos (código, nombre, precio, stock) y generación de "
        "reportes de stock en Excel y TXT.\n"
        "- Cierre de Caja: resume las ventas del día, calcula la ganancia estimada y permite exportar/imprimir el cierre.\n"
        "- Configuración: tema claro/oscuro, tamaño de letra, umbral de stock bajo, nombre del local, mensaje del "
        "ticket, alineación del encabezado, alertas de stock bajo por email y la clave de la API de Gemini que usa este asistente.\n"
        "- Asistente AI (vos): respondés preguntas sobre stock, ventas, ganancias, recomendaciones para vender más "
        "y cualquier duda sobre cómo usar la app.\n\n"
        f"RESUMEN ACTUAL DE LA APP:\n{resumen_app}\n\n"
        f"TOP PRODUCTOS MÁS VENDIDOS:\n{top_texto}\n\n"
        f"LISTADO COMPLETO DE STOCK ACTUAL:\n{lineas_stock}\n"
    )
    return contexto


def consultar_gemini(pregunta, contexto, cfg):
    """Llama a la API gratuita de Gemini (Google AI Studio) con el contexto de la app.
    Devuelve (texto_respuesta, error). error es None si la consulta fue exitosa."""
    api_key = (cfg or {}).get("gemini_api_key", "").strip()
    if not api_key:
        return None, "no_key"

    modelo = (cfg or {}).get("gemini_model", "gemini-2.5-flash").strip() or "gemini-2.5-flash"
    url = GEMINI_URL_TMPL.format(model=modelo)

    body = {
        "system_instruction": {"parts": [{"text": contexto}]},
        "contents": [{"role": "user", "parts": [{"text": pregunta}]}],
    }
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data,
        headers={"Content-Type": "application/json", "x-goog-api-key": api_key},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=25) as resp:
            resultado = json.loads(resp.read().decode("utf-8"))
        candidatos = resultado.get("candidates", [])
        if not candidatos:
            return None, "sin_respuesta"
        partes = candidatos[0].get("content", {}).get("parts", [])
        texto = "".join(p.get("text", "") for p in partes).strip()
        if not texto:
            return None, "sin_respuesta"
        return texto, None
    except urllib.error.HTTPError as e:
        try:
            detalle = e.read().decode("utf-8")
        except Exception:
            detalle = str(e)
        return None, f"Error {e.code} de la API de Gemini: {detalle}"
    except urllib.error.URLError as e:
        return None, f"Sin conexión a internet o servidor no disponible ({e.reason})."
    except Exception as e:
        return None, f"Error inesperado al consultar la IA: {e}"


def generar_respuesta_ai(pregunta, cfg=None):
    """Punto de entrada del asistente: intenta responder con Gemini (usando todo el
    contexto de la app) y, si no hay clave configurada o falla la conexión, usa el
    motor de respaldo basado en reglas para no dejar al usuario sin respuesta."""
    pregunta = str(pregunta or "").strip()
    if not pregunta:
        return (
            "Estoy listo para ayudarte con todo lo que esta aplicación sabe: stock, ventas, ganancias, reportes, alertas y configuración. "
            "Escribí lo que querés saber y te doy una respuesta basada en los datos existentes."
        )

    if cfg is not None:
        contexto = construir_contexto_app()
        respuesta, error = consultar_gemini(pregunta, contexto, cfg)
        if respuesta:
            return respuesta
        if error == "no_key":
            return (
                "⚠️ Todavía no configuraste la clave gratuita de la API de Gemini.\n"
                "Andá a Configuración → 'Clave de API de Gemini', pegá una clave gratuita obtenida en "
                "aistudio.google.com/apikey y guardá los cambios.\n\n"
                "Mientras tanto te respondo con el modo básico:\n" + _generar_respuesta_ai_basica(pregunta)
            )
        # Cualquier otro error (sin conexión, error HTTP, etc.): no dejamos al usuario sin respuesta.
        return (
            f"⚠️ No pude contactar a la IA en este momento ({error}). Te respondo con el modo básico:\n"
            + _generar_respuesta_ai_basica(pregunta)
        )

    return _generar_respuesta_ai_basica(pregunta)


def _generar_respuesta_ai_basica(pregunta):
    pregunta_norm = str(pregunta or "").strip().lower()
    ventas_hoy = obtener_ventas_del_dia()
    ventas_totales = obtener_ventas_generales()
    productos = listar_productos()
    top_vendidos = calcular_top_productos(ventas_totales, top_n=5)
    resumen_app = obtener_resumen_app(ventas_totales, productos)

    if any(p in pregunta_norm for p in ("cómo funciona", "que hace", "qué hace", "sabe la app", "todo de la app", "aplicación")):
        return (
            f"{CONOCIMIENTO_APP} {resumen_app} "
            "Puedo explicarte cómo usar cada función: desde cargar productos y ajustar stock, hasta cerrar caja, generar reportes y emitir tickets."
        )

    if "gráfica" in pregunta_norm or "grafica" in pregunta_norm or "top" in pregunta_norm:
        if plt is None:
            return (
                "Puedo ayudarte a generar una gráfica de los productos más vendidos, "
                "pero falta la librería matplotlib. Instalá con `pip install matplotlib`."
            )
        return (
            "Para ver los productos que más rotan, usaré los datos de ventas actuales. "
            "Presioná 'Generar gráfica Top Vendidos' y obtendrás un gráfico con los productos con mayor cantidad vendida."
        )

    if any(p in pregunta_norm for p in ("venta", "ventas", "vendido", "facturación", "ingresos", "total")):
        total = sum(v["subtotal"] for v in ventas_totales)
        unidades = sum(v["cantidad"] for v in ventas_totales)
        respuesta = [
            f"En el histórico hay {len(ventas_totales)} ventas registradas por un total de ${total:.2f}.",
            f"Se han vendido {unidades} unidades en total."
        ]
        if top_vendidos:
            respuesta.append("Los productos más vendidos son: " + ", ".join([f"{p[0]} ({p[1]['cantidad']} unidades)" for p in top_vendidos]))
        if ventas_hoy:
            hoy = sum(v["subtotal"] for v in ventas_hoy)
            respuesta.append(f"Hoy ya se vendió ${hoy:.2f}."
                             " Usá el cierre de caja para exportar el resumen del día.")
        return " ".join(respuesta)

    if any(p in pregunta_norm for p in ("ganancia", "ganancias", "margen", "beneficio")):
        ganancias = sum(v["ganancia"] for v in ventas_totales if v.get("ganancia") is not None)
        if ganancias > 0:
            return (
                f"La ganancia acumulada registrada es de ${ganancias:.2f}. "
                "Asegurate de ingresar el margen de cada venta en 'Nueva Venta' para mantener ese cálculo actualizado."
            )
        return (
            "Todavía no hay datos de ganancia claros en las ventas registradas. "
            "Usá el campo de margen en 'Nueva Venta' para calcularla automáticamente y luego revisala en el cierre de caja."
        )

    if any(p in pregunta_norm for p in ("stock", "inventario", "existencia", "cantidad")):
        productos_bajo = [p for p in productos if p[3] <= 3]
        respuesta = [f"Hay {len(productos)} productos registrados en stock."]
        if productos_bajo:
            respuesta.append(
                "Los productos con stock bajo son: "
                + ", ".join([f"{p[1]} ({p[3]} unidades)" for p in productos_bajo[:5]])
                + "."
            )
        else:
            respuesta.append("No hay productos críticos por debajo del umbral de stock bajo.")
        if top_vendidos:
            respuesta.append("Tus mejores productos actuales son: " + ", ".join([p[0] for p in top_vendidos[:3]]) + ".")
        return " ".join(respuesta)

    if any(p in pregunta_norm for p in ("recomend", "como vender", "qué debo hacer", "cómo puedo", "mejorar ventas", "generar más ventas", "mejorar")):
        return generar_recomendaciones(ventas_totales, productos)

    return (
        f"{CONOCIMIENTO_APP} {resumen_app} "
        "Puedo ayudarte con preguntas abiertas sobre cómo usar la app, cómo mejorar inventario y ventas, o qué datos revisar primero. "
        "Preguntá algo como '¿Qué producto debo promocionar?' o '¿Cómo reduzco productos sin venta?'."
    )


def generar_grafica_mas_vendidos(top_n=5):
    if not openpyxl or plt is None:
        return None, "No es posible generar la gráfica porque falta la librería matplotlib o los datos de Excel."
    ventas = obtener_ventas_generales()
    if not ventas:
        return None, "No hay datos de ventas para generar una gráfica."

    top = calcular_top_productos(ventas, top_n)
    if not top:
        return None, "No hay productos vendidos para crear una gráfica."

    nombres = [item[0] for item in top][::-1]
    cantidades = [item[1]["cantidad"] for item in top][::-1]
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.barh(nombres, cantidades, color=[LOGO_TEAL if i % 2 == 0 else LOGO_BLUE for i in range(len(nombres))])
    ax.set_title("Productos más vendidos")
    ax.set_xlabel("Unidades vendidas")
    ax.set_ylabel("Producto")
    ax.grid(axis="x", linestyle="--", alpha=0.3)
    for i, valor in enumerate(cantidades):
        ax.text(valor + 0.1, i, str(valor), va="center")

    path = os.path.join(BASE_DIR, "top_productos_mas_vendidos.png")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path, None


# ----------------------------------------------------------------------
# Reportes de stock (xlsx + txt)
# ----------------------------------------------------------------------
def generar_reportes_stock():
    productos = listar_productos()
    disponibles = [p for p in productos if p[3] > 0]
    sin_stock = [p for p in productos if p[3] <= 0]

    if openpyxl is not None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Stock"
        ws.append(["Código", "Nombre", "Precio", "Stock"])
        for p in productos:
            ws.append(list(p))
        wb.save(STOCK_XLSX)

    with open(STOCK_TXT, "w", encoding="utf-8") as f:
        f.write(f"Reporte de stock - {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write("=" * 50 + "\n\n")
        f.write("PRODUCTOS DISPONIBLES\n")
        f.write("-" * 50 + "\n")
        for p in disponibles:
            f.write(f"{p[1]}  |  ${p[2]:.2f}  |  Stock: {p[3]}\n")
        f.write("\nPRODUCTOS SIN STOCK\n")
        f.write("-" * 50 + "\n")
        for p in sin_stock:
            f.write(f"{p[1]}  |  ${p[2]:.2f}\n")

    return STOCK_XLSX, STOCK_TXT


# ----------------------------------------------------------------------
# Paletas de tema (colores tomados del logo)
# ----------------------------------------------------------------------
LOGO_TEAL = "#4dbf9e"       # verde azulado del logo
LOGO_BLUE = "#007877"       # azul profundo del logo
LOGO_NAVY = "#0a296c"       # azul navy del logo
LOGO_DARK = "#122235"       # esquina oscura del logo (fondo dark)

TEMAS = {
    "claro": {
        "bg": "#f4f7f8",            # fondo general, gris muy claro y neutro
        "fg": "#1b2733",            # texto principal, azul muy oscuro
        "panel": "#ffffff",         # tablas / paneles
        "panel_alt": "#f0f4f5",     # franja alterna sutil en listas
        "accent": LOGO_BLUE,        # botones y acentos
        "accent_hover": "#00615f",  # accent al pasar el mouse / presionar
        "accent_fg": "#ffffff",
        "accent2": LOGO_TEAL,       # segundo acento (para degradés y detalles)
        "borde": "#dde5e6",
        "muted": "#5b6b76",         # texto secundario
    },
    "oscuro": {
        "bg": LOGO_DARK,             # fondo general, navy oscuro del logo
        "fg": "#eaf5f2",             # texto claro
        "panel": "#1c3348",          # tablas / paneles
        "panel_alt": "#20395035",   # franja alterna sutil en listas (se ignora, ver nota abajo)
        "accent": LOGO_TEAL,         # botones y acentos (resalta sobre el fondo oscuro)
        "accent_hover": "#3aa686",
        "accent_fg": "#0a2018",
        "accent2": LOGO_BLUE,
        "borde": "#2a4a63",
        "muted": "#9db3c2",
    },
}
# nota: panel_alt de "oscuro" se recalcula abajo con un tono sólido válido
TEMAS["oscuro"]["panel_alt"] = "#213c53"

# Colores fijos para las alertas de stock (se ven igual en ambos temas para máxima claridad)
COLOR_STOCK_MEDIO_BG = "#fff3b0"     # fila amarilla: queda poco stock (4 a 7 unidades)
COLOR_STOCK_MEDIO_FG = "#5c4a00"
COLOR_STOCK_CRITICO_BG = "#f7b8b8"   # fila roja: stock crítico (0 a 3 unidades)
COLOR_STOCK_CRITICO_FG = "#6b0f0f"

GRADIENTE_LOBBY = (LOGO_TEAL, LOGO_NAVY)


def _hex_a_rgb(hexcolor):
    hexcolor = hexcolor.lstrip("#")
    return tuple(int(hexcolor[i:i + 2], 16) for i in (0, 2, 4))


def _interpolar_color(c1, c2, t):
    r1, g1, b1 = _hex_a_rgb(c1)
    r2, g2, b2 = _hex_a_rgb(c2)
    r = int(r1 + (r2 - r1) * t)
    g = int(g1 + (g2 - g1) * t)
    b = int(b1 + (b2 - b1) * t)
    return f"#{r:02x}{g:02x}{b:02x}"


class BotonDegradado(tk.Canvas):
    """Cuadrado con relleno en degradé (estilo del logo) que actúa como botón."""

    def __init__(self, parent, texto, comando, color_inicio, color_fin, fondo,
                 ancho=180, alto=180, fuente=("Segoe UI", 13, "bold")):
        super().__init__(parent, width=ancho, height=alto, highlightthickness=0, bg=fondo, bd=0)
        self.comando = comando
        self._dibujar_degradado(ancho, alto, color_inicio, color_fin)
        self.create_text(
            ancho / 2, alto / 2, text=texto, fill="#ffffff", font=fuente,
            width=ancho - 30, justify="center",
        )
        self.configure(cursor="hand2")
        self.bind("<Button-1>", lambda e: self.comando())

    def _dibujar_degradado(self, ancho, alto, color_inicio, color_fin):
        for i in range(ancho):
            t = i / max(1, (ancho - 1))
            color = _interpolar_color(color_inicio, color_fin, t)
            self.create_line(i, 0, i, alto, fill=color)
        self.create_rectangle(0, alto - 6, ancho, alto, fill=color_fin, outline=color_fin)


# ----------------------------------------------------------------------
# Aplicación principal
# ----------------------------------------------------------------------
class POSApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.cfg = cargar_config()
        self.title("Sistema de Punto de Venta")
        self.geometry("1000x680")
        self.minsize(880, 600)

        self.style = ttk.Style(self)
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass

        self.aplicar_tema()

        self.contenedor = tk.Frame(self, bg=self.paleta["bg"])
        self.contenedor.pack(fill="both", expand=True)

        # Contenedor de notificaciones: vive siempre arriba de todo (encima del lobby,
        # stock, etc.) y apila cada notificación automáticamente, sin cálculos manuales
        # de posición, para que nunca queden tapadas unas con otras.
        self._notif_container = tk.Frame(self, bg=self.paleta["bg"])
        self._notif_container.place(relx=0.5, y=14, anchor="n")

        self.frame_actual = None
        self.chat_window = None
        self.chat_historial = []
        self._avisos_desactivados = set()   # códigos de producto silenciados hasta que se repongan
        self._notif_por_producto = {}       # código -> widget de su notificación activa (máx. 1 por producto)
        self.mostrar_lobby()
        self.after(3000, self._chequear_stock_periodico)

    def aplicar_tema(self):
        self.paleta = TEMAS.get(self.cfg.get("tema", "claro"), TEMAS["claro"])
        tam = self.cfg.get("tamano_fuente", 11)
        self.fuente_normal = ("Segoe UI", tam)
        self.fuente_titulo = ("Segoe UI", tam + 8, "bold")
        self.fuente_boton = ("Segoe UI", tam + 1, "bold")
        self.configure(bg=self.paleta["bg"])

        self.style.configure("TFrame", background=self.paleta["bg"])
        self.style.configure("Panel.TFrame", background=self.paleta["panel"])
        self.style.configure("TLabel", background=self.paleta["bg"], foreground=self.paleta["fg"], font=self.fuente_normal)
        self.style.configure("Panel.TLabel", background=self.paleta["panel"], foreground=self.paleta["fg"], font=self.fuente_normal)
        self.style.configure("Titulo.TLabel", background=self.paleta["bg"], foreground=self.paleta["fg"], font=self.fuente_titulo)
        self.style.configure(
            "TButton", font=self.fuente_boton, padding=8, relief="flat",
            background=self.paleta["accent"], foreground=self.paleta["accent_fg"], borderwidth=0,
        )
        self.style.map(
            "TButton",
            background=[("active", self.paleta["accent_hover"]), ("pressed", self.paleta["accent_hover"])],
        )
        self.style.configure(
            "Treeview",
            background=self.paleta["panel"],
            fieldbackground=self.paleta["panel"],
            foreground=self.paleta["fg"],
            font=self.fuente_normal,
            rowheight=int(tam * 2.4),
            borderwidth=0,
            relief="flat",
        )
        self.style.configure(
            "Treeview.Heading", font=(self.fuente_normal[0], tam, "bold"),
            background=self.paleta["accent2"], foreground="#ffffff", relief="flat", padding=6,
        )
        self.style.map("Treeview.Heading", background=[("active", self.paleta["accent2"])])
        self.style.map("Treeview", background=[("selected", self.paleta["accent"])],
                        foreground=[("selected", self.paleta["accent_fg"])])
        self.style.configure("TEntry", font=self.fuente_normal, padding=4)
        self.style.configure("TCombobox", font=self.fuente_normal, padding=4)

    def refrescar_todo(self):
        self.aplicar_tema()
        self.contenedor.configure(bg=self.paleta["bg"])
        self._notif_container.configure(bg=self.paleta["bg"])
        self.mostrar_lobby()

    def cambiar_frame(self, frame_cls):
        if self.frame_actual is not None:
            self.frame_actual.destroy()
        self.frame_actual = frame_cls(self.contenedor, self)
        self.frame_actual.pack(fill="both", expand=True)

    def mostrar_lobby(self):
        self.cambiar_frame(LobbyFrame)

    def mostrar_nueva_venta(self):
        self.cambiar_frame(NuevaVentaFrame)

    def mostrar_stock(self):
        self.cambiar_frame(StockFrame)

    def mostrar_cierre_caja(self):
        self.cambiar_frame(CierreCajaFrame)

    def mostrar_configuracion(self):
        self.cambiar_frame(ConfiguracionFrame)

    def _crear_notificacion_producto(self, codigo_producto, mensaje, tipo):
        """Crea la notificación de UN producto puntual y la deja fija en pantalla
        (no se autocierra sola) hasta que el usuario la desactive o el producto se
        reponga. Nunca crea una segunda notificación para el mismo producto mientras
        ya tenga una activa."""
        if codigo_producto in self._notif_por_producto:
            return  # ya se está mostrando su notificación, no se duplica

        if tipo == "critico":
            bg, fg, icono = COLOR_STOCK_CRITICO_BG, COLOR_STOCK_CRITICO_FG, "🚨"
        else:
            bg, fg, icono = COLOR_STOCK_MEDIO_BG, COLOR_STOCK_MEDIO_FG, "⚠️"

        caja = tk.Frame(self._notif_container, bg=bg, highlightthickness=0, bd=0)
        tk.Label(
            caja, text=f"{icono}  {mensaje}", font=self.fuente_boton,
            bg=bg, fg=fg, padx=16, pady=10,
        ).pack(side="left")
        tk.Button(
            caja, text="Desactivar aviso", font=(self.fuente_normal[0], self.fuente_normal[1] - 1),
            bg=bg, fg=fg, relief="flat", bd=0, cursor="hand2",
            activebackground=bg, activeforeground=fg,
            # el código queda atado a ESTE botón puntual, no se comparte con otras notificaciones
            command=lambda c=codigo_producto: self._desactivar_aviso(c),
        ).pack(side="left", padx=(4, 16))

        caja.pack(pady=(0, 8))
        self._notif_container.lift()
        self._notif_por_producto[codigo_producto] = caja

    def _cerrar_notificacion_producto(self, codigo_producto):
        """Cierra (si existe) la notificación de un producto puntual, sin tocar
        las de los demás productos."""
        caja = self._notif_por_producto.pop(codigo_producto, None)
        if caja is not None:
            try:
                caja.destroy()
            except tk.TclError:
                pass

    def _cerrar_todas_notificaciones(self):
        """Cierra de una todas las notificaciones de stock que estén en pantalla
        (se usa al desactivar las notificaciones desde Configuración)."""
        for codigo in list(self._notif_por_producto.keys()):
            self._cerrar_notificacion_producto(codigo)

    def _desactivar_aviso(self, codigo_producto):
        """El usuario silencia los avisos de ESTE producto puntual (no afecta a los
        demás). Se reactiva solo apenas el stock de ese producto vuelva a estar
        por encima del mínimo (más de 7 unidades) y luego vuelva a bajar de nuevo."""
        self._avisos_desactivados.add(codigo_producto)
        self._cerrar_notificacion_producto(codigo_producto)

    def _chequear_stock_periodico(self):
        """Revisa el stock cada tantos minutos indique la Configuración (por
        defecto 1.5) y se ve en cualquier pantalla, incluido el lobby. Si las
        notificaciones están desactivadas, no se crea ningún aviso nuevo pero
        el chequeo se sigue reprogramando, para retomar apenas se reactiven.
        Por cada producto mantiene como máximo UNA notificación en pantalla,
        que queda fija hasta que se desactive o el producto se reponga. Un
        producto que vuelve a tener stock saludable (más de 7 unidades) cierra
        su notificación pendiente y olvida que estaba desactivado, para que si
        vuelve a bajar del mínimo te avise de nuevo con la opción de
        desactivar disponible otra vez."""
        try:
            intervalo_min = float(self.cfg.get("notificaciones_intervalo_minutos", 1.5))
            if intervalo_min <= 0:
                intervalo_min = 1.5
        except (TypeError, ValueError):
            intervalo_min = 1.5
        intervalo_ms = int(intervalo_min * 60 * 1000)

        if not self.cfg.get("notificaciones_activas", True):
            self.after(intervalo_ms, self._chequear_stock_periodico)
            return

        try:
            productos = listar_productos()
        except Exception:
            productos = []

        for codigo, nombre, precio, stock in productos:
            if stock > 7:
                # se repuso: se olvida el "desactivado" y se cierra cualquier aviso pendiente
                self._avisos_desactivados.discard(codigo)
                self._cerrar_notificacion_producto(codigo)
                continue

            if codigo in self._avisos_desactivados:
                continue

            if 0 <= stock <= 3:
                self._crear_notificacion_producto(codigo, f"'{nombre}' está en stock crítico", "critico")
            elif 4 <= stock <= 7:
                self._crear_notificacion_producto(codigo, f"Se está terminando el stock de '{nombre}'", "medio")

        self.after(intervalo_ms, self._chequear_stock_periodico)

    def abrir_asistente_ai(self):
        if self.chat_window is not None and self.chat_window.winfo_exists():
            self.chat_window.deiconify()
            self.chat_window.lift()
            self.chat_window.focus_force()
            return
        self.chat_window = AsistenteChatWindow(self)


# ----------------------------------------------------------------------
# Ventana del Asistente AI (chat aparte)
# ----------------------------------------------------------------------
class AsistenteChatWindow(tk.Toplevel):
    def __init__(self, app: POSApp):
        super().__init__(app)
        self.app = app
        self.title("Asistente AI")
        self.geometry("480x640")
        self.minsize(360, 420)
        self.configure(bg=app.paleta["panel"])
        self.protocol("WM_DELETE_WINDOW", self._cerrar)
        # Mientras es True, no se puede mandar una pregunta nueva: hay una
        # respuesta de la IA pendiente (ver enviar_pregunta / _mostrar_respuesta).
        self._esperando_respuesta = False

        header = tk.Frame(self, bg=app.paleta["panel"])
        header.pack(fill="x", padx=12, pady=(12, 4))
        tk.Label(
            header, text="🤖 Asistente AI",
            font=(app.fuente_titulo[0], app.fuente_normal[1] + 2, "bold"),
            bg=app.paleta["panel"], fg=app.paleta["fg"],
        ).pack(side="left")
        tk.Button(
            header, text="Limpiar chat", command=self._limpiar,
            font=app.fuente_boton, bg="#d94c4c", fg="#ffffff", relief="flat",
        ).pack(side="right")

        self.chat_text = tk.Text(
            self, wrap="word", font=app.fuente_normal,
            bg=app.paleta["bg"], fg=app.paleta["fg"], bd=1, relief="solid",
        )
        self.chat_text.pack(fill="both", expand=True, padx=12, pady=8)
        self.chat_text.tag_configure(
            "usuario", foreground=app.paleta["accent"],
            font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
        )
        self.chat_text.tag_configure(
            "ia_tag", foreground=app.paleta["accent2"],
            font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
        )
        self.chat_text.configure(state="disabled")

        entrada_frame = tk.Frame(self, bg=app.paleta["panel"])
        entrada_frame.pack(fill="x", padx=12, pady=(0, 12))
        self.entry_pregunta = tk.Entry(entrada_frame, font=app.fuente_normal)
        self.entry_pregunta.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.entry_pregunta.bind("<Return>", self.enviar_pregunta)
        self.btn_enviar = tk.Button(
            entrada_frame, text="Enviar", command=self.enviar_pregunta,
            font=app.fuente_boton, bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat",
        )
        self.btn_enviar.pack(side="left")

        self._cargar_historial()
        self.entry_pregunta.focus_set()

    def _cargar_historial(self):
        self.chat_text.configure(state="normal")
        self.chat_text.delete("1.0", tk.END)
        if not self.app.chat_historial:
            self.chat_text.insert(
                tk.END,
                "Preguntame lo que quieras sobre stock, ventas, ganancias o cómo usar la app.\n\n",
            )
        for pregunta, respuesta in self.app.chat_historial:
            self.chat_text.insert(tk.END, "Vos: ", "usuario")
            self.chat_text.insert(tk.END, pregunta + "\n")
            self.chat_text.insert(tk.END, "IA: ", "ia_tag")
            self.chat_text.insert(tk.END, respuesta + "\n\n")
        self.chat_text.configure(state="disabled")
        self.chat_text.see(tk.END)

    def enviar_pregunta(self, event=None):
        if self._esperando_respuesta:
            # Ya hay una pregunta en curso: no se manda otra hasta que
            # termine de responder (con Enter tampoco, no solo con el botón).
            return
        pregunta = self.entry_pregunta.get().strip()
        if not pregunta:
            return
        self.entry_pregunta.delete(0, tk.END)
        self._esperando_respuesta = True
        self.entry_pregunta.configure(state="disabled")
        self.btn_enviar.configure(state="disabled")

        self.chat_text.configure(state="normal")
        self.chat_text.insert(tk.END, "Vos: ", "usuario")
        self.chat_text.insert(tk.END, pregunta + "\n")
        self.chat_text.insert(tk.END, "IA: ", "ia_tag")
        self._respuesta_start = self.chat_text.index(tk.END)
        self.chat_text.insert(tk.END, "Pensando... ⏳\n\n")
        self.chat_text.configure(state="disabled")
        self.chat_text.see(tk.END)

        def _consultar():
            try:
                respuesta = generar_respuesta_ai(pregunta, self.app.cfg)
            except Exception as e:
                respuesta = f"⚠️ Ocurrió un error inesperado al consultar la IA: {e}"
            self.after(0, lambda: self._mostrar_respuesta(pregunta, respuesta))

        threading.Thread(target=_consultar, daemon=True).start()

    def _mostrar_respuesta(self, pregunta, respuesta):
        self._esperando_respuesta = False
        try:
            self.chat_text.configure(state="normal")
            self.chat_text.delete(self._respuesta_start, tk.END)
            self.chat_text.insert(tk.END, respuesta + "\n\n")
            self.chat_text.configure(state="disabled")
            self.chat_text.see(tk.END)
            self.app.chat_historial.append((pregunta, respuesta))
            self.entry_pregunta.configure(state="normal")
            self.btn_enviar.configure(state="normal")
            self.entry_pregunta.focus_set()
        except tk.TclError:
            pass  # la ventana se cerró mientras se esperaba la respuesta

    def _limpiar(self):
        self.app.chat_historial.clear()
        self._cargar_historial()

    def _cerrar(self):
        self.app.chat_window = None
        self.destroy()


# ----------------------------------------------------------------------
# Lobby
# ----------------------------------------------------------------------
class LobbyFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self._logo_ref = None

        if HAS_PIL and os.path.exists(LOGO_PATH):
            try:
                img = Image.open(LOGO_PATH).convert("RGBA").resize((96, 96), Image.LANCZOS)
                self._logo_ref = ImageTk.PhotoImage(img)
                tk.Label(self, image=self._logo_ref, bg=app.paleta["bg"]).pack(pady=(35, 5))
            except Exception:
                pass

        tk.Label(
            self, text="Punto de Venta", font=app.fuente_titulo,
            bg=app.paleta["bg"], fg=app.paleta["fg"]
        ).pack(pady=(5 if self._logo_ref else 50, 30))

        botones_frame = tk.Frame(self, bg=app.paleta["bg"])
        botones_frame.pack(fill="both", expand=True, padx=20, pady=10)

        color_inicio, color_fin = GRADIENTE_LOBBY
        opciones = [
            ("Nueva\nVenta", app.mostrar_nueva_venta),
            ("Administrar\nStock", app.mostrar_stock),
            ("Cierre de\nCaja", app.mostrar_cierre_caja),
            ("Configuración", app.mostrar_configuracion),
            ("🤖 Asistente\nAI", app.abrir_asistente_ai),
        ]

        for i, (texto, comando) in enumerate(opciones):
            cuadro = BotonDegradado(
                botones_frame, texto, comando, color_inicio, color_fin,
                fondo=app.paleta["bg"], fuente=(app.fuente_titulo[0], app.fuente_normal[1] + 3, "bold"),
            )
            cuadro.grid(row=0, column=i, padx=12, pady=20)


# ----------------------------------------------------------------------
# Cierre de Caja
# ----------------------------------------------------------------------
class CierreCajaFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self.ventas_hoy = obtener_ventas_del_dia()

        # Top bar
        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(
            top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal,
            bg=app.paleta["panel"], fg=app.paleta["accent"], relief="flat", bd=0,
            padx=12, pady=5, cursor="hand2", activebackground=app.paleta["borde"],
        ).pack(side="left")
        tk.Label(top, text="Cierre de Caja del Día", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        # Tabla de ventas del día
        tabla_frame = tk.Frame(self, bg=app.paleta["bg"])
        tabla_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("hora", "producto", "precio", "cantidad", "subtotal")
        self.tree = ttk.Treeview(tabla_frame, columns=columnas, show="headings", selectmode="browse")
        for col, texto, ancho in [
            ("hora", "Hora", 90), ("producto", "Producto", 280),
            ("precio", "Precio Unit.", 110), ("cantidad", "Cant.", 80),
            ("subtotal", "Subtotal", 120)
        ]:
            self.tree.heading(col, text=texto)
            self.tree.column(col, width=ancho, anchor="center")

        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(tabla_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)

        # Totales
        totales_frame = tk.Frame(self, bg=app.paleta["bg"])
        totales_frame.pack(fill="x", padx=20, pady=10)

        self.total_ganado = sum(v["subtotal"] for v in self.ventas_hoy)
        
        # Ganancia estimada = suma de la ganancia calculada con el margen (%)
        # cargado en "Nueva venta" al finalizar cada compra.
        self.ganancia_estimada = 0.0
        for v in self.ventas_hoy:
            if v.get("ganancia") is not None:
                self.ganancia_estimada += v["ganancia"]
            else:
                # Ventas viejas guardadas antes de tener columna "Ganancia"
                prod = buscar_producto_por_nombre(v["producto"])
                if prod:
                    costo_base = prod[2]
                    subtotal_costo = costo_base * v["cantidad"]
                    self.ganancia_estimada += max(0.0, v["subtotal"] - subtotal_costo)
                else:
                    self.ganancia_estimada += v["subtotal"] * 0.20

        self.lbl_total_ganado = tk.Label(
            totales_frame, text=f"Total Ganado Hoy (Ventas): $ {self.total_ganado:.2f}",
            font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["accent"]
        )
        self.lbl_total_ganado.pack(anchor="w")

        self.lbl_ganancia_neta = tk.Label(
            totales_frame, text=f"Ganancia Estimada Hoy: $ {self.ganancia_estimada:.2f}",
            font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]
        )
        self.lbl_ganancia_neta.pack(anchor="w", pady=(5, 0))

        # Botones de Acción
        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))

        tk.Button(acciones_frame, text="💾 Guardar Cierre en Excel", command=self.guardar_excel, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🖨 Imprimir (Térmica)", command=self.imprimir_termica, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🖨 Imprimir (Normal)", command=self.imprimir_normal, font=app.fuente_boton,
                  bg=LOGO_TEAL, fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)

        self.cargar_tabla()

    def cargar_tabla(self):
        self.tree.delete(*self.tree.get_children())
        for v in self.ventas_hoy:
            self.tree.insert("", "end", values=(
                v["hora"], v["producto"], f"{v['precio_unitario']:.2f}", v["cantidad"], f"{v['subtotal']:.2f}"
            ))

    def guardar_excel(self):
        if not openpyxl:
            messagebox.showwarning("Excel no disponible", "No se encontró el módulo 'openpyxl'.")
            return
        if not self.ventas_hoy:
            messagebox.showinfo("Sin datos", "No hay ventas registradas el día de hoy para exportar.")
            return

        nombre_archivo = f"cierre_caja_{datetime.now().strftime('%Y%m%d')}.xlsx"
        path_cierre = os.path.join(BASE_DIR, nombre_archivo)

        wb = Workbook()
        ws = wb.active
        ws.title = "Cierre de Caja"
        ws.append(["Hora", "Producto", "Precio Unitario", "Cantidad", "Subtotal"])

        for v in self.ventas_hoy:
            ws.append([v["hora"], v["producto"], v["precio_unitario"], v["cantidad"], v["subtotal"]])

        ws.append([])
        ws.append(["TOTAL VENDIDO HOY:", self.total_ganado])
        ws.append(["GANANCIA ESTIMADA HOY:", self.ganancia_estimada])

        wb.save(path_cierre)
        messagebox.showinfo("Cierre Guardado", f"El informe de cierre de caja se guardó en:\n{path_cierre}")

    def imprimir_termica(self):
        if not self.ventas_hoy:
            messagebox.showinfo("Sin ventas", "No hay ventas registradas hoy para imprimir.")
            return

        items_cierre = [{
            "codigo": "",
            "nombre": v["producto"],
            "precio": v["precio_unitario"],
            "cantidad": v["cantidad"]
        } for v in self.ventas_hoy]

        imprimir_ticket(self.app.cfg, items_cierre, self.total_ganado, 0.0, self.total_ganado)

    def imprimir_normal(self):
        if not self.ventas_hoy:
            messagebox.showinfo("Sin ventas", "No hay ventas registradas hoy para imprimir.")
            return

        fecha_str = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        lineas = []
        lineas.append(f"=========================================")
        lineas.append(f"       REPORTE DE CIERRE DE CAJA        ")
        lineas.append(f" Local: {self.app.cfg.get('nombre_local', 'MI COMERCIO')}")
        lineas.append(f" Fecha: {fecha_str}")
        lineas.append(f"=========================================")
        lineas.append(f"{'Hora':<8} {'Producto':<20} {'Cant.':<6} {'Subtotal':<10}")
        lineas.append(f"-----------------------------------------")
        for v in self.ventas_hoy:
            lineas.append(f"{v['hora']:<8} {v['producto'][:18]:<20} {v['cantidad']:<6} ${v['subtotal']:.2f}")
        lineas.append(f"-----------------------------------------")
        lineas.append(f"TOTAL VENDIDO HOY:      $ {self.total_ganado:.2f}")
        lineas.append(f"GANANCIA ESTIMADA HOY:  $ {self.ganancia_estimada:.2f}")
        lineas.append(f"=========================================")
        lineas.append(f" {self.app.cfg.get('mensaje_ticket', '¡Gracias por su compra!')}")

        texto = "\n".join(lineas)
        if imprimir_normal_texto(texto, titulo_documento="Cierre_de_Caja"):
            messagebox.showinfo("Impresión", "Documento enviado a la impresora predeterminada.")


# ----------------------------------------------------------------------
# Nueva Venta
# ----------------------------------------------------------------------
class NuevaVentaFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app
        self.items = {}
        # Estado para detectar automáticamente un escaneo de código de barras
        # (que escribe todo el código en pocos milisegundos) sin necesidad de
        # que el usuario toque Enter. Ver _on_tecla_codigo / _evaluar_auto_escaneo.
        self._inicio_tecleo = None
        self._ultimo_tecleo = None
        self._auto_scan_job = None

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(
            top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal,
            bg=app.paleta["panel"], fg=app.paleta["accent"], relief="flat", bd=0,
            padx=12, pady=5, cursor="hand2", activebackground=app.paleta["borde"],
        ).pack(side="left")
        tk.Label(top, text="Nueva Venta", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        scan_frame = tk.Frame(self, bg=app.paleta["bg"])
        scan_frame.pack(fill="x", padx=20, pady=10)
        tk.Label(scan_frame, text="Escanear código de barras:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left")
        self.entry_codigo = tk.Entry(scan_frame, font=app.fuente_normal, width=30)
        self.entry_codigo.pack(side="left", padx=10)
        self.entry_codigo.bind("<Return>", self.escanear)
        self.entry_codigo.bind("<KeyRelease>", self._on_tecla_codigo)
        self.entry_codigo.focus_set()
        tk.Button(scan_frame, text="Agregar", command=self.escanear, font=app.fuente_normal).pack(side="left")

        tabla_frame = tk.Frame(self, bg=app.paleta["bg"])
        tabla_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("codigo", "nombre", "precio", "cantidad", "subtotal")
        self.tree = ttk.Treeview(tabla_frame, columns=columnas, show="headings", selectmode="browse")
        for col, texto, ancho in [
            ("codigo", "Código", 120), ("nombre", "Producto", 260),
            ("precio", "Precio unit.", 110), ("cantidad", "Cantidad", 90),
            ("subtotal", "Subtotal", 110),
        ]:
            self.tree.heading(col, text=texto)
            self.tree.column(col, width=ancho, anchor="center")
        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(tabla_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.bind("<Double-1>", self.editar_cantidad)

        totales_frame = tk.Frame(self, bg=app.paleta["bg"])
        totales_frame.pack(fill="x", padx=20, pady=10)

        tk.Label(totales_frame, text="Margen de ganancia (%):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=0, column=0, sticky="w")
        self.entry_margen = tk.Entry(totales_frame, font=app.fuente_normal, width=8)
        self.entry_margen.insert(0, "0")
        self.entry_margen.grid(row=0, column=1, padx=10, sticky="w")
        self.entry_margen.bind("<KeyRelease>", lambda e: self.recalcular())

        self.lbl_subtotal = tk.Label(totales_frame, text="Subtotal: $ 0.00", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"])
        self.lbl_subtotal.grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))
        self.lbl_total = tk.Label(totales_frame, text="TOTAL: $ 0.00", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["accent"])
        self.lbl_total.grid(row=2, column=0, columnspan=2, sticky="w")

        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))

        tk.Button(acciones_frame, text="🖨 Imprimir ticket", command=self.imprimir, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="💾 Guardar en Excel", command=self.guardar_excel, font=app.fuente_boton,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="✔ Finalizar compra", command=self.finalizar, font=app.fuente_boton,
                  bg=LOGO_TEAL, fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)
        tk.Button(acciones_frame, text="🗑 Eliminar producto", command=self.eliminar_producto, font=app.fuente_boton,
                  bg="#d94c4c", fg="#ffffff", relief="flat").pack(side="left", padx=5, ipadx=10, ipady=6)

    def escanear(self, event=None):
        if self._auto_scan_job is not None:
            self.after_cancel(self._auto_scan_job)
            self._auto_scan_job = None
        self._inicio_tecleo = None
        self._ultimo_tecleo = None

        codigo = self.entry_codigo.get().strip()
        self.entry_codigo.delete(0, tk.END)
        if not codigo:
            return
        producto = buscar_producto(codigo)
        if producto is None:
            messagebox.showwarning("No encontrado", f"El código '{codigo}' no está registrado en Stock.")
            return
        _, nombre, precio, stock = producto
        if codigo in self.items:
            self.items[codigo]["cantidad"] += 1
        else:
            self.items[codigo] = {"codigo": codigo, "nombre": nombre, "precio": precio, "cantidad": 1}
        self.refrescar_tabla()

    def _on_tecla_codigo(self, event):
        """Se dispara con cada tecla que se suelta en el campo de código. Un
        lector de código de barras escribe el código entero en pocos
        milisegundos; si detecta esa velocidad, agrega el producto solo, sin
        que haga falta tocar Enter. Si en cambio se escribe a mano, a ritmo
        humano, no hace nada distinto: se sigue confirmando con Enter o el
        botón 'Agregar', como antes."""
        if self._auto_scan_job is not None:
            self.after_cancel(self._auto_scan_job)
            self._auto_scan_job = None

        texto = self.entry_codigo.get()
        if not texto:
            self._inicio_tecleo = None
            self._ultimo_tecleo = None
            return

        ahora = time.monotonic()
        if len(texto) <= 1:
            self._inicio_tecleo = ahora
        self._ultimo_tecleo = ahora

        # Esperamos una breve pausa sin nuevas teclas antes de decidir: así
        # nos aseguramos de evaluar el código ya completo, no a mitad de camino.
        self._auto_scan_job = self.after(120, self._evaluar_auto_escaneo)

    def _evaluar_auto_escaneo(self):
        self._auto_scan_job = None
        texto = self.entry_codigo.get().strip()
        if len(texto) < 3 or self._inicio_tecleo is None or self._ultimo_tecleo is None:
            return

        duracion_total = max(0.0, self._ultimo_tecleo - self._inicio_tecleo)
        intervalos = max(1, len(texto) - 1)
        promedio_por_caracter = duracion_total / intervalos

        # Un lector de código de barras es muchísimo más rápido que cualquier
        # persona tipeando (todo el código entra casi de una); si el promedio
        # por caracter es bajísimo, lo tratamos como escaneo y se agrega solo.
        if promedio_por_caracter <= 0.04:
            self._inicio_tecleo = None
            self._ultimo_tecleo = None
            self.escanear()

    def refrescar_tabla(self):
        self.tree.delete(*self.tree.get_children())
        for it in self.items.values():
            subtotal = it["precio"] * it["cantidad"]
            self.tree.insert("", "end", iid=it["codigo"], values=(
                it["codigo"], it["nombre"], f"{it['precio']:.2f}", it["cantidad"], f"{subtotal:.2f}",
            ))
        self.recalcular()

    def recalcular(self):
        subtotal = sum(it["precio"] * it["cantidad"] for it in self.items.values())
        try:
            margen = float(self.entry_margen.get().replace(",", "."))
        except ValueError:
            margen = 0.0
        total = subtotal * (1 + margen / 100)
        self.lbl_subtotal.config(text=f"Subtotal: $ {subtotal:.2f}")
        self.lbl_total.config(text=f"TOTAL: $ {total:.2f}")
        return subtotal, margen, total

    def editar_cantidad(self, event):
        item_id = self.tree.identify_row(event.y)
        columna = self.tree.identify_column(event.x)
        if not item_id or columna != "#4":
            return
        actual = self.items[item_id]["cantidad"]
        nueva = simpledialog.askinteger("Editar cantidad", f"Cantidad para '{self.items[item_id]['nombre']}':",
                                         initialvalue=actual, minvalue=1)
        if nueva:
            self.items[item_id]["cantidad"] = nueva
            self.refrescar_tabla()

    def eliminar_producto(self):
        seleccion = self.tree.selection()
        if not seleccion:
            messagebox.showinfo("Eliminar producto", "Seleccioná un producto de la lista primero.")
            return
        codigo = seleccion[0]
        item = self.items[codigo]
        if item["cantidad"] > 1:
            cantidad_a_borrar = simpledialog.askinteger(
                "Eliminar producto",
                f"'{item['nombre']}' tiene {item['cantidad']} unidades cargadas.\n"
                f"¿Cuántas unidades querés quitar?",
                initialvalue=item["cantidad"], minvalue=1, maxvalue=item["cantidad"],
            )
            if not cantidad_a_borrar:
                return
            if cantidad_a_borrar >= item["cantidad"]:
                del self.items[codigo]
            else:
                item["cantidad"] -= cantidad_a_borrar
        else:
            del self.items[codigo]
        self.refrescar_tabla()

    def _items_como_lista(self):
        return list(self.items.values())

    def imprimir(self):
        if not self.items:
            messagebox.showinfo("Ticket vacío", "Todavía no escaneaste ningún producto.")
            return
        subtotal, margen, total = self.recalcular()
        imprimir_ticket(self.app.cfg, self._items_como_lista(), subtotal, margen, total)

    def guardar_excel(self):
        if not self.items:
            messagebox.showinfo("Ticket vacío", "Todavía no escaneaste ningún producto.")
            return
        _, margen, _ = self.recalcular()
        if guardar_venta_excel(self._items_como_lista(), margen):
            messagebox.showinfo("Guardado", f"La venta se guardó en:\n{VENTAS_XLSX}")

    def finalizar(self):
        if not self.items:
            messagebox.showinfo("Venta vacía", "Todavía no escaneaste ningún producto.")
            return
        _, margen, _ = self.recalcular()
        guardar_venta_excel(self._items_como_lista(), margen)
        umbral = self.app.cfg.get("umbral_stock_bajo", 3)
        for it in self.items.values():
            nombre, restante = descontar_stock(it["codigo"], it["cantidad"])
            if nombre is not None and restante <= umbral:
                enviar_alerta_stock(self.app.cfg, nombre, restante)
        messagebox.showinfo("Compra finalizada", "La venta se registró y el stock fue actualizado.")
        self.app.mostrar_lobby()


# ----------------------------------------------------------------------
# Administrar Stock
# ----------------------------------------------------------------------
class StockFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(
            top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal,
            bg=app.paleta["panel"], fg=app.paleta["accent"], relief="flat", bd=0,
            padx=12, pady=5, cursor="hand2", activebackground=app.paleta["borde"],
        ).pack(side="left")
        tk.Label(top, text="Administrar Stock", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        alta_frame = tk.LabelFrame(self, text="Registrar / actualizar producto", font=app.fuente_normal,
                                    bg=app.paleta["bg"], fg=app.paleta["fg"])
        alta_frame.pack(fill="x", padx=20, pady=10)

        tk.Label(alta_frame, text="Código de barras:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=0, column=0, padx=5, pady=8, sticky="e")
        self.entry_codigo = tk.Entry(alta_frame, font=app.fuente_normal, width=20)
        self.entry_codigo.grid(row=0, column=1, padx=5)
        self.entry_codigo.bind("<Return>", self.buscar_para_editar)

        tk.Label(alta_frame, text="Nombre:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=0, column=2, padx=5, sticky="e")
        self.entry_nombre = tk.Entry(alta_frame, font=app.fuente_normal, width=25)
        self.entry_nombre.grid(row=0, column=3, padx=5)

        tk.Label(alta_frame, text="Precio:", bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=1, column=0, padx=5, sticky="e")
        self.entry_precio = tk.Entry(alta_frame, font=app.fuente_normal, width=20)
        self.entry_precio.grid(row=1, column=1, padx=5)

        self.var_label_stock = tk.StringVar(value="Stock inicial:")
        tk.Label(alta_frame, textvariable=self.var_label_stock, bg=app.paleta["bg"], fg=app.paleta["fg"], font=app.fuente_normal).grid(row=1, column=2, padx=5, sticky="e")
        self.entry_stock = tk.Entry(alta_frame, font=app.fuente_normal, width=25)
        self.entry_stock.grid(row=1, column=3, padx=5)
        self.entry_stock.bind("<Return>", lambda e: self.guardar())

        tk.Button(alta_frame, text="Guardar producto", command=self.guardar, font=app.fuente_normal).grid(row=0, column=4, rowspan=2, padx=15)

        self.lbl_info = tk.Label(alta_frame, text="", font=(app.fuente_normal[0], app.fuente_normal[1] - 1),
                                  bg=app.paleta["bg"], fg=app.paleta["accent"], justify="left")
        self.lbl_info.grid(row=2, column=0, columnspan=5, sticky="w", padx=5, pady=(0, 8))

        lista_frame = tk.Frame(self, bg=app.paleta["bg"])
        lista_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columnas = ("nombre", "precio", "stock", "codigo")
        self.tree = ttk.Treeview(lista_frame, columns=columnas, show="headings", selectmode="browse")
        self.tree.heading("nombre", text="Producto")
        self.tree.heading("precio", text="Precio")
        self.tree.heading("stock", text="Stock")
        self.tree.heading("codigo", text="Código")
        self.tree.column("nombre", width=320)
        self.tree.column("precio", width=100, anchor="center")
        self.tree.column("stock", width=100, anchor="center")
        self.tree.column("codigo", width=140, anchor="center")
        self.tree.pack(fill="both", expand=True, side="left")
        scroll = ttk.Scrollbar(lista_frame, orient="vertical", command=self.tree.yview)
        scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.bind("<<TreeviewSelect>>", self.cargar_seleccion)

        acciones_frame = tk.Frame(self, bg=app.paleta["bg"])
        acciones_frame.pack(fill="x", padx=20, pady=(5, 20))
        tk.Button(acciones_frame, text="Eliminar producto seleccionado", command=self.eliminar, font=app.fuente_normal).pack(side="left", padx=5)
        tk.Button(acciones_frame, text="📄 Generar reporte de stock (Excel + TXT)", command=self.generar_reporte, font=app.fuente_normal).pack(side="left", padx=5)

        self.cargar_lista()

    def cargar_lista(self):
        self.tree.delete(*self.tree.get_children())
        self.tree.tag_configure("stock_medio", background=COLOR_STOCK_MEDIO_BG, foreground=COLOR_STOCK_MEDIO_FG)
        self.tree.tag_configure("stock_critico", background=COLOR_STOCK_CRITICO_BG, foreground=COLOR_STOCK_CRITICO_FG)

        for codigo, nombre, precio, stock in listar_productos():
            tags = ()
            if 0 <= stock <= 3:
                tags = ("stock_critico",)
            elif 4 <= stock <= 7:
                tags = ("stock_medio",)
            self.tree.insert("", "end", iid=codigo, values=(nombre, f"{precio:.2f}", stock, codigo), tags=tags)

    def buscar_para_editar(self, event=None):
        codigo = self.entry_codigo.get().strip()
        if not codigo:
            return
        self._cargar_en_formulario(codigo)
        self.entry_stock.focus_set()

    def cargar_seleccion(self, event=None):
        seleccion = self.tree.selection()
        if not seleccion:
            return
        codigo = seleccion[0]
        self.entry_codigo.delete(0, tk.END)
        self.entry_codigo.insert(0, codigo)
        self._cargar_en_formulario(codigo)

    def _cargar_en_formulario(self, codigo):
        producto = buscar_producto(codigo)
        self.entry_nombre.delete(0, tk.END)
        self.entry_precio.delete(0, tk.END)
        self.entry_stock.delete(0, tk.END)
        if producto:
            _, nombre, precio, stock = producto
            self.entry_nombre.insert(0, nombre)
            self.entry_precio.insert(0, str(precio))
            self.var_label_stock.set("Cantidad a agregar:")
            self.lbl_info.config(
                text=f"✔ Producto ya registrado — stock actual: {stock} unidades. "
                     f"La cantidad que cargues acá se SUMARÁ a ese stock (dejalo vacío para no tocar el stock).",
            )
        else:
            self.var_label_stock.set("Stock inicial:")
            self.lbl_info.config(text="🆕 Producto nuevo — completá nombre, precio y stock inicial.")
            self.entry_nombre.focus_set()

    def guardar(self):
        codigo = self.entry_codigo.get().strip()
        nombre = self.entry_nombre.get().strip()
        precio_txt = self.entry_precio.get().strip().replace(",", ".")
        stock_txt = self.entry_stock.get().strip()
        if not codigo or not nombre or not precio_txt:
            messagebox.showwarning("Datos incompletos", "Completá código, nombre y precio.")
            return
        try:
            precio = float(precio_txt)
        except ValueError:
            messagebox.showwarning("Datos inválidos", "El precio debe ser numérico.")
            return

        cantidad_ingresada = 0
        if stock_txt:
            try:
                cantidad_ingresada = int(stock_txt)
            except ValueError:
                messagebox.showwarning("Datos inválidos", "La cantidad/stock debe ser un número entero.")
                return

        producto_existente = buscar_producto(codigo)
        if producto_existente:
            stock_actual = producto_existente[3]
            nuevo_stock = stock_actual + cantidad_ingresada
            guardar_producto(codigo, nombre, precio, nuevo_stock)
            if cantidad_ingresada:
                messagebox.showinfo(
                    "Stock actualizado",
                    f"'{nombre}': {stock_actual} + {cantidad_ingresada} = {nuevo_stock} unidades.",
                )
        else:
            if not stock_txt:
                messagebox.showwarning("Datos incompletos", "Para un producto nuevo indicá el stock inicial.")
                return
            guardar_producto(codigo, nombre, precio, cantidad_ingresada)

        self.cargar_lista()
        self.entry_codigo.delete(0, tk.END)
        self.entry_nombre.delete(0, tk.END)
        self.entry_precio.delete(0, tk.END)
        self.entry_stock.delete(0, tk.END)
        self.var_label_stock.set("Stock inicial:")
        self.lbl_info.config(text="")
        self.entry_codigo.focus_set()

    def eliminar(self):
        seleccion = self.tree.selection()
        if not seleccion:
            messagebox.showinfo("Eliminar producto", "Seleccioná un producto de la lista primero.")
            return
        codigo = seleccion[0]
        if messagebox.askyesno("Confirmar", "¿Eliminar este producto del stock?"):
            eliminar_producto(codigo)
            self.cargar_lista()

    def generar_reporte(self):
        xlsx_path, txt_path = generar_reportes_stock()
        messagebox.showinfo("Reporte generado", f"Se generaron:\n{xlsx_path}\n{txt_path}")


# ----------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------
class ConfiguracionFrame(tk.Frame):
    def __init__(self, parent, app: POSApp):
        super().__init__(parent, bg=app.paleta["bg"])
        self.app = app

        top = tk.Frame(self, bg=app.paleta["bg"])
        top.pack(fill="x", padx=20, pady=(15, 5))
        tk.Button(
            top, text="\u2190 Volver", command=app.mostrar_lobby, font=app.fuente_normal,
            bg=app.paleta["panel"], fg=app.paleta["accent"], relief="flat", bd=0,
            padx=12, pady=5, cursor="hand2", activebackground=app.paleta["borde"],
        ).pack(side="left")
        tk.Label(top, text="Configuración", font=app.fuente_titulo, bg=app.paleta["bg"], fg=app.paleta["fg"]).pack(side="left", padx=20)

        # --- Área con scroll: los campos viven dentro de un Canvas para poder
        # desplazarse verticalmente cuando no entran todos en la ventana
        # (rueda del mouse, touchpad o la barra lateral).
        scroll_contenedor = tk.Frame(self, bg=app.paleta["bg"])
        scroll_contenedor.pack(fill="both", expand=True, padx=30, pady=10)

        canvas = tk.Canvas(scroll_contenedor, bg=app.paleta["bg"], highlightthickness=0, bd=0)
        scrollbar = ttk.Scrollbar(scroll_contenedor, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        panel = tk.Frame(canvas, bg=app.paleta["bg"])
        panel_id = canvas.create_window((0, 0), window=panel, anchor="nw")

        def _actualizar_scrollregion(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
        panel.bind("<Configure>", _actualizar_scrollregion)

        def _ajustar_ancho_panel(event):
            canvas.itemconfig(panel_id, width=event.width)
        canvas.bind("<Configure>", _ajustar_ancho_panel)

        # Rueda del mouse (Windows/Mac) y botones 4/5 (Linux), activa solo
        # mientras el cursor está sobre esta pantalla para no interferir con
        # el scroll de otras partes de la app.
        def _on_mousewheel(event):
            if getattr(event, "num", None) == 4:
                canvas.yview_scroll(-1, "units")
            elif getattr(event, "num", None) == 5:
                canvas.yview_scroll(1, "units")
            elif getattr(event, "delta", 0):
                canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")

        def _activar_scroll(event):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)
            canvas.bind_all("<Button-4>", _on_mousewheel)
            canvas.bind_all("<Button-5>", _on_mousewheel)

        def _desactivar_scroll(event):
            canvas.unbind_all("<MouseWheel>")
            canvas.unbind_all("<Button-4>")
            canvas.unbind_all("<Button-5>")

        canvas.bind("<Enter>", _activar_scroll)
        canvas.bind("<Leave>", _desactivar_scroll)

        # Tema
        tk.Label(panel, text="Tema:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=0, column=0, sticky="w", pady=5)
        self.var_tema = tk.StringVar(value=app.cfg.get("tema", "claro"))
        tk.Radiobutton(panel, text="Claro", variable=self.var_tema, value="claro", font=app.fuente_normal,
                        bg=app.paleta["bg"], fg=app.paleta["fg"], selectcolor=app.paleta["panel"]).grid(row=0, column=1, sticky="w")
        tk.Radiobutton(panel, text="Oscuro", variable=self.var_tema, value="oscuro", font=app.fuente_normal,
                        bg=app.paleta["bg"], fg=app.paleta["fg"], selectcolor=app.paleta["panel"]).grid(row=0, column=2, sticky="w")

        # Tamaño de letra
        tk.Label(panel, text="Tamaño de letra:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=1, column=0, sticky="w", pady=5)
        self.var_tamano = tk.IntVar(value=app.cfg.get("tamano_fuente", 11))
        tk.Spinbox(panel, from_=8, to=20, textvariable=self.var_tamano, width=5, font=app.fuente_normal).grid(row=1, column=1, sticky="w")

        # Umbral de stock bajo
        tk.Label(panel, text="Umbral de stock bajo (unidades):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=2, column=0, sticky="w", pady=5)
        self.var_umbral = tk.IntVar(value=app.cfg.get("umbral_stock_bajo", 3))
        tk.Spinbox(panel, from_=1, to=100, textvariable=self.var_umbral, width=5, font=app.fuente_normal).grid(row=2, column=1, sticky="w")

        # --- Edición de Texto de Impresión (Petición) ---
        tk.Label(panel, text="Personalización de Impresión", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=3, column=0, columnspan=3, sticky="w", pady=(15, 5))

        tk.Label(panel, text="Nombre del Local:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=4, column=0, sticky="w", pady=5)
        self.entry_nombre_local = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_nombre_local.insert(0, app.cfg.get("nombre_local", "MI COMERCIO"))
        self.entry_nombre_local.grid(row=4, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Mensaje en Ticket:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=5, column=0, sticky="w", pady=5)
        self.entry_mensaje_ticket = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_mensaje_ticket.insert(0, app.cfg.get("mensaje_ticket", "¡Gracias por su compra!"))
        self.entry_mensaje_ticket.grid(row=5, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Alineación del Encabezado:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=6, column=0, sticky="w", pady=5)
        self.var_disposicion = tk.StringVar(value=app.cfg.get("disposicion_ticket", "Centro"))
        combo_disp = ttk.Combobox(panel, textvariable=self.var_disposicion, values=["Izquierda", "Centro", "Derecha"], state="readonly", width=15)
        combo_disp.grid(row=6, column=1, sticky="w")

        # --- Alertas por email ---
        tk.Label(panel, text="Avisos de stock bajo por email", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=7, column=0, columnspan=3, sticky="w", pady=(15, 5))

        tk.Label(panel, text="Email que recibe el aviso:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=8, column=0, sticky="w", pady=5)
        self.entry_email_alertas = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_email_alertas.insert(0, app.cfg.get("email_alertas", ""))
        self.entry_email_alertas.grid(row=8, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Cuenta Gmail que envía (remitente):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=9, column=0, sticky="w", pady=5)
        self.entry_email_remitente = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_email_remitente.insert(0, app.cfg.get("email_remitente", ""))
        self.entry_email_remitente.grid(row=9, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Contraseña de aplicación de Gmail:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=10, column=0, sticky="w", pady=5)
        self.entry_email_password = tk.Entry(panel, font=app.fuente_normal, width=35, show="*")
        self.entry_email_password.insert(0, app.cfg.get("email_password", ""))
        self.entry_email_password.grid(row=10, column=1, columnspan=2, sticky="w")

        # --- Impresora ---
        tk.Label(panel, text="Nombre exacto de la impresora (opcional):", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=11, column=0, sticky="w", pady=5)
        self.entry_impresora = tk.Entry(panel, font=app.fuente_normal, width=35)
        self.entry_impresora.insert(0, app.cfg.get("nombre_impresora", ""))
        self.entry_impresora.grid(row=11, column=1, columnspan=2, sticky="w")

        # --- Asistente AI (Gemini) ---
        tk.Label(panel, text="Asistente AI (Gemini)", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=12, column=0, columnspan=3, sticky="w", pady=(15, 5))

        tk.Label(panel, text="Clave de API de Gemini:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=13, column=0, sticky="w", pady=5)
        self.entry_gemini_key = tk.Entry(panel, font=app.fuente_normal, width=35, show="*")
        self.entry_gemini_key.insert(0, app.cfg.get("gemini_api_key", ""))
        self.entry_gemini_key.grid(row=13, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Conseguí una clave gratis en aistudio.google.com/apikey",
                 font=(app.fuente_normal[0], app.fuente_normal[1] - 1), bg=app.paleta["bg"], fg=app.paleta["fg"]
                 ).grid(row=14, column=1, columnspan=2, sticky="w")

        tk.Label(panel, text="Modelo de Gemini:", font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=15, column=0, sticky="w", pady=5)
        self.entry_gemini_model = tk.Entry(panel, font=app.fuente_normal, width=25)
        self.entry_gemini_model.insert(0, app.cfg.get("gemini_model", "gemini-2.5-flash"))
        self.entry_gemini_model.grid(row=15, column=1, sticky="w")

        # --- Notificaciones en pantalla ---
        tk.Label(panel, text="Notificaciones en pantalla", font=(app.fuente_normal[0], app.fuente_normal[1], "bold"),
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=16, column=0, columnspan=3, sticky="w", pady=(15, 5))

        self.var_notif_activas = tk.BooleanVar(value=app.cfg.get("notificaciones_activas", True))
        tk.Checkbutton(
            panel, text="Activar notificaciones de stock bajo", variable=self.var_notif_activas,
            font=app.fuente_normal, bg=app.paleta["bg"], fg=app.paleta["fg"], selectcolor=app.paleta["panel"],
            command=lambda: self.spin_notif_intervalo.config(
                state="normal" if self.var_notif_activas.get() else "disabled"
            ),
        ).grid(row=17, column=0, columnspan=3, sticky="w", pady=5)

        tk.Label(panel, text="Frecuencia de las notificaciones (minutos):", font=app.fuente_normal,
                 bg=app.paleta["bg"], fg=app.paleta["fg"]).grid(row=18, column=0, sticky="w", pady=5)
        self.var_notif_intervalo = tk.DoubleVar(value=float(app.cfg.get("notificaciones_intervalo_minutos", 1.5)))
        self.spin_notif_intervalo = tk.Spinbox(
            panel, from_=0.5, to=60, increment=0.5, textvariable=self.var_notif_intervalo,
            width=6, font=app.fuente_normal, format="%.1f",
            state=("normal" if self.var_notif_activas.get() else "disabled"),
        )
        self.spin_notif_intervalo.grid(row=18, column=1, sticky="w")

        # --- Guardar ---
        tk.Button(panel, text="Guardar configuración", font=app.fuente_boton, command=self.guardar,
                  bg=app.paleta["accent"], fg=app.paleta["accent_fg"], relief="flat").grid(row=19, column=0, pady=15, ipadx=10, ipady=6, sticky="w")

    def guardar(self):
        self.app.cfg["tema"] = self.var_tema.get()
        self.app.cfg["tamano_fuente"] = int(self.var_tamano.get())
        self.app.cfg["umbral_stock_bajo"] = int(self.var_umbral.get())
        self.app.cfg["nombre_local"] = self.entry_nombre_local.get().strip()
        self.app.cfg["mensaje_ticket"] = self.entry_mensaje_ticket.get().strip()
        self.app.cfg["disposicion_ticket"] = self.var_disposicion.get()
        self.app.cfg["email_alertas"] = self.entry_email_alertas.get().strip()
        self.app.cfg["email_remitente"] = self.entry_email_remitente.get().strip()
        self.app.cfg["email_password"] = self.entry_email_password.get().strip()
        self.app.cfg["nombre_impresora"] = self.entry_impresora.get().strip()
        self.app.cfg["gemini_api_key"] = self.entry_gemini_key.get().strip()
        self.app.cfg["gemini_model"] = self.entry_gemini_model.get().strip() or "gemini-2.5-flash"

        notif_activas = bool(self.var_notif_activas.get())
        try:
            notif_intervalo = float(self.var_notif_intervalo.get())
        except (tk.TclError, ValueError):
            notif_intervalo = 1.5
        if notif_intervalo <= 0:
            notif_intervalo = 1.5
        self.app.cfg["notificaciones_activas"] = notif_activas
        self.app.cfg["notificaciones_intervalo_minutos"] = notif_intervalo

        guardar_config(self.app.cfg)
        if not notif_activas:
            self.app._cerrar_todas_notificaciones()
        messagebox.showinfo("Configuración", "Los cambios se guardaron correctamente.")
        self.app.refrescar_todo()


# ----------------------------------------------------------------------
# Punto de entrada
# ----------------------------------------------------------------------
def main():
    app = POSApp()
    app.mainloop()


if __name__ == "__main__":
    main()